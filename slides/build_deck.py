"""Dựng slide báo cáo từ template DUT bằng python-pptx."""
import copy
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu
from lxml import etree

TEMPLATE, OUT = sys.argv[1], sys.argv[2]
FIG = Path("E:/DataScience/src/figures")
FIGS = FIG / "slides"   # bản hình cho slide (chữ ≥ 12pt)

BLUE = RGBColor(0x07, 0x52, 0x89)       # xanh Bách Khoa (nền template)
BLUE2 = RGBColor(0x17, 0x5D, 0x98)
YELLOW = RGBColor(0xFF, 0xD4, 0x00)     # vàng đậm, đọc được trên nền trắng khi làm khối
ORANGE = RGBColor(0xEB, 0x68, 0x34)     # nhấn (cùng màu cam trong biểu đồ notebook)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x67, 0x75)
TINT = RGBColor(0xEA, 0xF2, 0xFA)       # nền thẻ
TINT_Y = RGBColor(0xFF, 0xF6, 0xCC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation(TEMPLATE)
tpl = list(prs.slides)
S_TITLE, S_SECTION, S_THANKS = tpl[0], tpl[3], tpl[14]
L_CONTENT = prs.slide_masters[1].slide_layouts[0]      # nền trắng, tòa nhà mờ, sóng vàng, logo
L_SECTION = S_SECTION.slide_layout

order = []          # thứ tự slide cuối cùng


# ---------------------------------------------------------------- tiện ích chung
def strip_placeholders(slide):
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)


def add_slide_number(slide):
    # Thêm placeholder số trang (kế thừa vị trí, kiểu chữ từ layout)
    xml = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:nvSpPr><p:cNvPr id="900" name="Slide Number"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        '<p:nvPr><p:ph type="sldNum" sz="quarter" idx="12"/></p:nvPr></p:nvSpPr><p:spPr/>'
        '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum">'
        '<a:rPr lang="vi-VN"/><a:t>‹#›</a:t></a:fld><a:endParaRPr lang="vi-VN"/></a:p></p:txBody></p:sp>')
    slide.shapes._spTree.append(etree.fromstring(xml))


def style_run(run, size, color=INK, bold=False, italic=False, font=FONT):
    f = run.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "vi-VN")


def set_bullet(p, color=BLUE, indent=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr")).set("val", str(color))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, space_after=6, line=1.1, name=None):
    """paras: list; mỗi phần tử là str hoặc list các (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            style_run(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", False))
        if bullets:
            set_bullet(p)
    return tb


def box(slide, x, y, w, h, fill=TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def badge(slide, x, y, d, label, fill=BLUE, color=WHITE, size=16):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    style_run(r, size, color, bold=True)
    return c


def arrow(slide, x, y, w, h, fill=YELLOW):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = fill
    a.line.fill.background(); a.shadow.inherit = False
    return a


def picture(slide, path, x, y, w, h, align="center"):
    """Chèn ảnh vừa khung (giữ tỉ lệ)."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else x
    py = y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    return pic


def table(slide, x, y, w, rows, col_w, size=12, row_h=0.36, header_fill=BLUE, first_bold=True):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"     # No Style, Table Grid
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else TINT
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            runs = [(val, {})] if isinstance(val, str) else val
            for t, o in runs:
                r = p.add_run(); r.text = t
                if i == 0:
                    style_run(r, size, WHITE, bold=True)
                else:
                    style_run(r, o.get("size", size), o.get("color", INK),
                              bold=o.get("bold", first_bold and j == 0))
            # viền mảnh xám nhạt
            tcPr = cell._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(side), w="6350")
                sf = etree.SubElement(ln, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr")).set("val", "C9D6E3")
    return gt


def style_chart(chart, legend=True, size=12):
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = MUTED
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE8, 0xEE)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(size - 1)
    ca = chart.category_axis
    ca.format.line.color.rgb = RGBColor(0xBF, 0xC7, 0xD1)
    ca.tick_labels.font.size = Pt(size)


def content(kicker, title, notes=""):
    s = prs.slides.add_slide(L_CONTENT)
    strip_placeholders(s)
    text(s, 0.6, 0.3, 10.5, 0.35, [kicker.upper()], size=12, color=ORANGE, bold=True, name="Kicker")
    text(s, 0.6, 0.6, 12.1, 0.8, [title], size=28, color=BLUE, bold=True, name="Title")
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    order.append(s)
    return s


def set_para(p, new_text, size=None):
    """Giữ định dạng run đầu tiên của đoạn, thay toàn bộ chữ."""
    runs = p.runs
    if not runs:
        r = p.add_run(); r.text = new_text; return
    runs[0].text = new_text
    if size:
        runs[0].font.size = Pt(size)
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def section(num, title, notes=""):
    s = prs.slides.add_slide(L_SECTION)
    strip_placeholders(s)
    src_csld = S_SECTION._element.find(qn("p:cSld"))
    bg = src_csld.find(qn("p:bg"))
    if bg is not None:
        s._element.find(qn("p:cSld")).insert(0, copy.deepcopy(bg))
    for el in S_SECTION.shapes._spTree.iterchildren():
        if el.tag in (qn("p:sp"), qn("p:cxnSp"), qn("p:pic"), qn("p:grpSp")):
            s.shapes._spTree.append(copy.deepcopy(el))
    tb = [sh for sh in s.shapes if sh.has_text_frame and "PHẦN" in sh.text_frame.text][0]
    tb.width = Inches(5.3)
    ps = tb.text_frame.paragraphs
    set_para(ps[0], f"PHẦN {num}")
    set_para(ps[1], title, size=38)
    for sh in s.shapes:
        if sh.shape_id != tb.shape_id and not sh.has_text_frame or (sh.shape_id != tb.shape_id and not sh.text_frame.text.strip()):
            sh.left = tb.left + tb.width + Inches(0.2)
    for p in ps[2:]:
        p._p.getparent().remove(p._p)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    order.append(s)
    return s


# ================================================================ NỘI DUNG 41 SLIDE
# Không dùng slide ngăn cách phần: mỗi slide có dòng nhãn nhỏ (kicker) cho biết đang ở phần nào.
GRAY = RGBColor(0x9A, 0xA0, 0xA6)


def caption(slide, x, y, w, t, size=12):
    text(slide, x, y, w, 0.4, [t], size=size, color=MUTED)


def callout(slide, x, y, w, h, title, body, dark=False, size=13):
    box(slide, x, y, w, h, fill=BLUE if dark else TINT)
    text(slide, x + 0.2, y + 0.08, w - 0.4, 0.4, [title], size=size + 2, color=YELLOW if dark else BLUE, bold=True)
    text(slide, x + 0.2, y + 0.5, w - 0.4, h - 0.55, body if isinstance(body, list) else [body], size=size,
         color=WHITE if dark else INK, space_after=4)


# ---------------------------------------------------------------- 1. Bìa
ps = [sh for sh in S_TITLE.shapes if sh.has_text_frame]
main = max(ps, key=lambda sh: len(sh.text_frame.paragraphs))
paras = main.text_frame.paragraphs
set_para(paras[0], "Đề tài môn Khoa học dữ liệu")
set_para(paras[1], "Dự đoán giá ô tô cũ trên Chợ Tốt bằng Random Forest", size=34)
for shp in ps:
    for p in shp.text_frame.paragraphs:
        t = p.text.strip()
        if t.startswith("SVTH"):
            set_para(p, "SVTH: <Họ tên SV 1> – <Họ tên SV 2>")
        elif t.upper().startswith("KHO"):
            set_para(p, "GVHD: <Tên giảng viên>")
S_TITLE.notes_slide.notes_text_frame.text = (
    "Chào thầy cô. Nhóm em trình bày đề tài dự đoán giá rao bán ô tô trên Chợ Tốt. Mô hình được phân tích kỹ là Random Forest, "
    "làm hai bản: dùng thư viện scikit-learn và tự cài đặt bằng NumPy. Ridge Regression chỉ dùng làm mốc so sánh.")
order.append(S_TITLE)

# ---------------------------------------------------------------- 2. Nội dung
s = content("Nội dung trình bày", "Năm phần của bài báo cáo",
            "Bài trình bày đi theo đúng quy trình một dự án học máy từ đầu đến cuối.")
items = [("01", "Giới thiệu đề tài", "Bài toán, mục tiêu và quy trình thực hiện"),
         ("02", "Thu thập và khám phá dữ liệu", "Crawl 14.835 tin rao bán ô tô từ Chợ Tốt"),
         ("03", "Tiền xử lý và xử lý dữ liệu", "Quy tắc cứng; chia train/test rồi học mọi tham số chỉ từ train"),
         ("04", "Mô hình Random Forest", "Lý thuyết, bản thư viện, bản tự cài đặt, kiểm chứng; Ridge để so sánh"),
         ("05", "Kết quả và kết luận", "Đánh giá trên tập test, diễn giải, thí nghiệm, hạn chế")]
for k, (n, t, d) in enumerate(items):
    y = 1.6 + k * 0.88
    badge(s, 1.2, y, 0.68, n, fill=BLUE if k % 2 == 0 else ORANGE, size=18)
    text(s, 2.1, y - 0.02, 9.5, 0.4, [t], size=20, color=INK, bold=True)
    text(s, 2.1, y + 0.36, 9.5, 0.35, [d], size=14, color=MUTED)

# ---------------------------------------------------------------- 3. Bài toán
s = content("Phần 1 · Giới thiệu", "Bài toán: ước lượng giá một chiếc xe từ thông tin trong tin đăng",
            "Người mua và người bán xe cũ đều cần một mốc giá hợp lý. Đây là bài toán hồi quy có giám sát: đầu vào là các thuộc tính "
            "ai cũng thấy trong tin đăng, đầu ra là giá rao bán tính bằng triệu đồng.")
box(s, 0.6, 1.65, 5.0, 4.25, fill=TINT)
text(s, 0.85, 1.8, 4.6, 0.4, ["Đầu vào: thông tin trong tin đăng"], size=16, color=BLUE, bold=True)
text(s, 0.85, 2.3, 4.6, 3.5, ["Hãng, dòng xe", "Năm sản xuất, số km đã đi", "Hộp số, nhiên liệu",
                              "Kiểu dáng, số chỗ", "Xuất xứ, tình trạng (mới / đã dùng)"],
     size=16, bullets=True, space_after=8)
arrow(s, 5.85, 3.45, 0.9, 0.6)
box(s, 7.0, 1.65, 5.7, 1.95, fill=BLUE)
text(s, 7.25, 1.8, 5.2, 0.4, ["Đầu ra"], size=16, color=YELLOW, bold=True)
text(s, 7.25, 2.2, 5.2, 1.3, [[("Giá rao bán", {"size": 26, "bold": True, "color": WHITE})],
                               [("(triệu đồng) - bài toán hồi quy có giám sát", {"size": 15, "color": WHITE})]])
box(s, 7.0, 3.85, 5.7, 2.05, fill=TINT_Y)
text(s, 7.25, 3.98, 5.2, 0.4, ["Yêu cầu của đề tài"], size=16, color=BLUE, bold=True)
text(s, 7.25, 4.4, 5.3, 1.5, ["Một mô hình từ đầu đến cuối, có thước đo đánh giá",
                              "Phân tích kỹ Random Forest, tự cài đặt và kiểm chứng",
                              "Có kỹ thuật xử lý dữ liệu mất cân bằng"],
     size=14, bullets=True, space_after=5)

# ---------------------------------------------------------------- 4. Quy trình
s = content("Phần 1 · Giới thiệu", "Quy trình thực hiện",
            "Năm bước, mỗi bước một notebook. Tiền xử lý chỉ dùng quy tắc cứng; mọi bước học thống kê từ dữ liệu (điền thiếu, gộp nhãn, lọc nhãn sai) "
            "làm sau khi chia train/test và chỉ học từ train. Tập test chỉ dùng một lần ở bước đánh giá.")
steps = [("Thu thập", "Crawl Chợ Tốt\n14.835 tin"), ("Tiền xử lý", "Quy tắc cứng\n14.066 tin"),
         ("Xử lý dữ liệu", "Chia tập, fit trên train\n11.153 / 2.797"), ("Mô hình", "Random Forest\nthư viện + tự cài đặt"),
         ("Đánh giá", "MAPE, R², ±20%\ntheo phân khúc")]
for k, (t, d) in enumerate(steps):
    x = 0.6 + k * 2.5
    box(s, x, 1.8, 2.1, 2.6, fill=TINT if k % 2 == 0 else TINT_Y)
    badge(s, x + 0.7, 2.0, 0.7, str(k + 1), fill=BLUE if k % 2 == 0 else ORANGE, size=20)
    text(s, x + 0.1, 2.85, 1.9, 0.45, [t], size=18, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.1, 3.35, 1.9, 0.95, d.split("\n"), size=13, color=INK, align=PP_ALIGN.CENTER, space_after=2)
    if k < 4:
        arrow(s, x + 2.13, 2.9, 0.32, 0.4)
text(s, 0.6, 4.7, 12.1, 0.5, [[("Chống rò rỉ dữ liệu: ", {"bold": True, "color": BLUE}),
                              ("chia tập trước; điền thiếu, gộp nhãn, mã hóa, tham số chỉ học từ train, rồi áp nguyên lên test.", {})]], size=15)
text(s, 0.6, 5.3, 12.1, 0.5, [[("Tái lập: ", {"bold": True, "color": BLUE}),
                              ("mọi bước ngẫu nhiên cố định random_state = 42, chạy lại cho đúng cùng kết quả.", {})]], size=15)

# ---------------------------------------------------------------- 5. Thu thập
IMG = Path("E:/DataScience/image")
s = content("Phần 2 · Thu thập dữ liệu", "Crawl 14.835 tin rao bán ô tô từ Chợ Tốt",
            "Bước 1 (ảnh trái): duyệt 750 trang danh sách mới nhất mua-ban-oto?page=1..750, lấy link từng tin, bỏ link trùng (tin ưu tiên lặp lại ở mọi trang). "
            "Bước 2 (ảnh phải): vào trang chi tiết, đọc các thẻ HTML có thuộc tính itemprop; khung cam là thẻ itemprop='mileage_v2' chứa số km 44.000. "
            "Dùng requests + BeautifulSoup, nghỉ 0,5 giây mỗi tin và 1 giây mỗi trang, tự lùi thời gian khi bị trả mã 403/429, chạy tiếp được khi bị ngắt. "
            "Chỉ lấy thông tin công khai, không lấy thông tin liên hệ người bán (tên và số điện thoại trong ảnh đã được che). "
            "Khung xanh: trang còn các mục số đời chủ, còn hạn đăng kiểm mà code chưa lấy; đây là hướng phát triển.")
text(s, 0.6, 1.4, 12.1, 0.4, [[("750 trang  ·  14.835 tin  ·  12 thuộc tính mỗi tin  ·  thu thập ngày 25/10/2026", {"bold": True, "color": ORANGE})]], size=16)


def shot(path, x, y, w, h, marks, label, n):
    pic = picture(s, path, x, y, w, h)
    iw, ih = Image.open(path).size
    sc = pic.width / Emu(Inches(1)) / iw if False else (pic.width / 914400) / iw
    px, py = pic.left / 914400, pic.top / 914400
    for (x0, y0, x1, y1), col in marks:
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(px + x0 * sc), Inches(py + y0 * sc),
                               Inches((x1 - x0) * sc), Inches((y1 - y0) * sc))
        r.fill.background(); r.line.color.rgb = col; r.line.width = Pt(2.5); r.shadow.inherit = False
    badge(s, px - 0.15, py - 0.15, 0.45, str(n), fill=ORANGE if n == 2 else BLUE, size=14)
    text(s, px + 0.4, 1.83, pic.width / 914400 - 0.4, 0.4, [label], size=14, color=BLUE, bold=True)


shot(IMG / "Screenshot_1_che.png", 0.8, 2.3, 4.9, 3.65, [], "Trang danh sách: lấy link từng tin", 1)
arrow(s, 5.95, 3.9, 0.7, 0.45)
shot(IMG / "Screenshot_3_che.png", 6.9, 2.3, 4.9, 3.65,
     [((723, 352, 1235, 390), ORANGE), ((352, 312, 570, 412), BLUE)],
     "Trang chi tiết: đọc thẻ itemprop", 2)
text(s, 11.9, 2.4, 1.2, 3.0, [[("khung cam:", {"bold": True, "color": ORANGE})], "thẻ itemprop số km",
                              [("khung xanh:", {"bold": True, "color": BLUE})], "số đời chủ, đăng kiểm: chưa lấy"],
     size=11, space_after=3)

# ---------------------------------------------------------------- 6. Từ điển dữ liệu
s = content("Phần 2 · Thu thập dữ liệu", "Từ điển dữ liệu: biết cột nào người bán gõ tự do",
            "Cột chọn từ danh sách thì sạch, không lỗi chính tả. Hai cột gõ tự do là số km và giá, cũng là hai nguồn lỗi chính "
            "cần kiểm tra ở bước tiền xử lý. URL chứa vị trí tin đăng.")
rows = [["Cột", "Ý nghĩa", "Người bán nhập", "Xử lý"],
        ["carbrand / carmodel", "Hãng / dòng xe", "chọn từ danh sách", "gộp nhãn hiếm theo cấp"],
        ["mfdate", "Năm sản xuất", "chọn (có \"trước 1980\")", "→ tuổi xe = 2026 − năm"],
        ["mileage_v2", "Số km đã đi", [("gõ tự do", {"color": ORANGE, "bold": True})], "kiểm tra km vô lý, điền thiếu"],
        ["gearbox / fuel", "Hộp số / nhiên liệu", "chọn (3 / 4 giá trị)", "one-hot"],
        ["cartype / carseats", "Kiểu dáng / số chỗ", "chọn, không bắt buộc", "điền theo dòng xe"],
        ["carorigin", "Xuất xứ", "chọn, không bắt buộc", "thiếu 41% → \"Không rõ\""],
        ["condition_ad", "Tình trạng", "Mới / Đã sử dụng", "sửa xe \"Mới\" quá 3 năm tuổi"],
        ["veh_warranty_policy", "Bảo hành", "một giá trị duy nhất", "bỏ cột"],
        ["price", "Giá rao bán (đồng)", [("gõ tự do", {"color": ORANGE, "bold": True})], "biến mục tiêu, loại giá sai"],
        ["url", "Đường dẫn tin", "tự sinh, chứa tỉnh/thành", "trích tỉnh/thành rồi bỏ"]]
table(s, 0.6, 1.55, 12.1, rows, [2.7, 2.7, 3.1, 3.6], size=13, row_h=0.41)

# ---------------------------------------------------------------- 7. EDA: giá trị thiếu
s = content("Phần 2 · Khám phá dữ liệu", "Dữ liệu thô: thiếu tập trung ở 4 cột không bắt buộc",
            "Trên 13 đặc trưng, chỉ 4 cột có thiếu đáng kể: xuất xứ 40,5%, kiểu dáng 32,9%, số chỗ 24,9%, số km 12,2%; giá chỉ thiếu 1 tin. "
            "8 cột còn lại đầy đủ vì là ô bắt buộc. Hình phải tách theo tình trạng xe: xe mới thiếu số km 100% vì trang xe mới không có ô km, "
            "xe đã dùng chỉ 1,8%, nên đây là thiếu có cấu trúc, điền 0 km. Chỉ 37% số tin đầy đủ mọi ô, nên phải điền chứ không xóa dòng.")
picture(s, FIG / "v01_thieu_moi_dac_trung.png", 0.45, 1.45, 6.3, 3.2)
picture(s, FIG / "v03_thieu_theo_tinh_trang.png", 6.9, 1.5, 6.0, 3.0)
caption(s, 0.55, 4.65, 6.1, "Thiếu trên cả 13 đặc trưng (14.835 tin thô), chia theo nguồn thiếu", 13)
caption(s, 7.0, 4.65, 5.8, "Tỉ lệ thiếu ở xe mới so với xe đã dùng", 13)
cards = [("40,5%", "xuất xứ thiếu, không suy được → nhãn \"Không rõ\""),
         ("100%", "xe mới thiếu số km (xe cũ 1,8%) → thiếu có cấu trúc"),
         ("37%", "số tin đủ mọi ô → điền thiếu, không xóa dòng")]
for k, (v, d) in enumerate(cards):
    x = 0.6 + k * 4.1
    box(s, x, 5.15, 3.8, 1.0, fill=TINT)
    text(s, x + 0.12, 5.15, 1.4, 1.0, [v], size=24, color=ORANGE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.55, 5.15, 2.15, 1.0, [d], size=12.5, anchor=MSO_ANCHOR.MIDDLE, space_after=0)

# ---------------------------------------------------------------- 7b. EDA: cơ cấu kiểu dáng, số chỗ, xuất xứ
s = content("Phần 2 · Khám phá dữ liệu", "Kiểu dáng, số chỗ, xuất xứ: lệch mạnh và thiếu nhiều",
            "Mỗi thanh là phần trăm số tin của một giá trị, thanh xám là phần thiếu. Kiểu dáng: SUV và Sedan chiếm phần lớn, "
            "Coupe và Mui trần dưới 0,5%. Số chỗ: gần như chỉ 5 chỗ và 7 chỗ. Xuất xứ: phần thiếu 40,5% còn lớn hơn cả Việt Nam 28,8%. "
            "Kiểu dáng và số chỗ do dòng xe quyết định nên điền theo dòng xe; xuất xứ không suy được nên để nhãn Không rõ.")
picture(s, FIG / "v04_co_cau_kieu_dang_so_cho_xuat_xu.png", 0.4, 1.45, 9.1, 3.75)
caption(s, 0.5, 5.3, 8.9, "Cơ cấu trên 14.835 tin thô; thanh xám = thiếu (chưa khai)", 13)
side = [("Lệch mạnh", "SUV + Sedan ≈ 44%; 5 và 7 chỗ ≈ 63%. Coupe, Mui trần, 10-12 chỗ dưới 0,5%: mô hình có rất ít mẫu."),
        ("Thiếu là nhóm lớn", "Thiếu kiểu dáng 32,9%, xuất xứ 40,5%: lớn hơn mọi nhãn thật, không thể bỏ qua."),
        ("Hướng xử lý", "Kiểu dáng, số chỗ: điền theo dòng xe. Xuất xứ: nhãn \"Không rõ\" (slide 15, 18).")]
for k, (t, d) in enumerate(side):
    callout(s, 9.75, 1.5 + k * 1.55, 3.2, 1.4, t, d, size=12)

# ---------------------------------------------------------------- 8. EDA: giá
s = content("Phần 2 · Khám phá dữ liệu", "Tuổi xe và thương hiệu quyết định phần lớn giá",
            "Giá trung vị gần như đứng yên 3 năm đầu (khoảng 600-620 triệu) rồi giảm đều: 498 triệu ở 5 năm, 335 ở 10 năm, 205 ở 15 năm. "
            "Trên thang log, giảm đều nghĩa là mất giá theo tỉ lệ % gần như cố định mỗi năm, nên học trên log giá. "
            "Ba hãng sang Lexus, BMW, Mercedes có giá trung vị gấp 2-4 lần nhóm phổ thông.")
picture(s, FIGS / "s01_eda_gia_theo_tuoi.png", 0.5, 1.5, 6.1, 3.7)
picture(s, FIGS / "s02_eda_gia_theo_hang.png", 6.75, 1.5, 6.1, 3.7)
caption(s, 0.6, 5.3, 5.9, "Trung vị: xe mới khoảng 620 triệu, 10 năm 335, 15 năm 205 triệu", 14)
caption(s, 6.85, 5.3, 5.9, "Lexus 1.820, BMW 1.000, Mercedes 900; Toyota 409, Chevrolet 145 triệu", 14)

# ---------------------------------------------------------------- 8b. EDA: tương quan, kiểu dáng, hộp số
s = content("Phần 2 · Khám phá dữ liệu", "Tuổi xe tương quan mạnh nhất; kiểu dáng và hộp số tách giá rõ",
            "Tương quan Pearson với log giá: tuổi xe −0,60, số km −0,38, số chỗ 0,09. Tuổi và số km tương quan 0,54 nên mang trùng một phần "
            "thông tin; vì vậy khi đọc độ quan trọng đặc trưng phải kiểm tra thêm bằng permutation. Giá trung vị: SUV 600 triệu, bán tải 565, "
            "Sedan 387, Hatchback 250; số sàn 210 so với tự động 515 (xe số sàn phần lớn là xe cũ).")
picture(s, FIGS / "s04_eda_tuong_quan.png", 0.4, 1.45, 4.6, 3.75)
picture(s, FIGS / "s05_eda_gia_theo_kieu_dang.png", 5.1, 1.5, 7.8, 3.6)
cards = [("−0,60", "tương quan tuổi xe với log giá: mạnh nhất"),
         ("0,54", "tuổi và số km trùng thông tin → cần permutation importance"),
         ("2,5×", "xe số tự động đắt hơn số sàn (515 so với 210 triệu)")]
for k, (v, d) in enumerate(cards):
    x = 0.6 + k * 4.1
    box(s, x, 5.3, 3.8, 0.85, fill=TINT)
    text(s, x + 0.12, 5.3, 1.3, 0.85, [v], size=22, color=ORANGE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.45, 5.3, 2.25, 0.85, [d], size=12, anchor=MSO_ANCHOR.MIDDLE, space_after=0)

# ---------------------------------------------------------------- 9. Cơ chế thu thập → hệ quả cho tiền xử lý
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Cách dữ liệu được tạo ra quyết định cách xử lý",
            "Hiểu code crawl giúp biết mỗi loại lỗi đến từ đâu. Ví dụ N/A là do code không tìm thấy thẻ, tức người bán bỏ trống, "
            "còn 'Khác' là lựa chọn thật của người bán nên phải giữ.")
rows = [["Cơ chế", "Hệ quả", "Cách xử lý"],
        ["Không tìm thấy thẻ → ghi \"N/A\"", "Thiếu thật", "đổi về NaN (đọc với keep_default_na=False)"],
        ["Trang hiện \"Đang cập nhật\" khi bỏ trống xuất xứ", "Cũng là thiếu thật", "đổi về NaN"],
        ["Danh sách có sẵn \"Khác\", \"Dòng khác\"", "Lựa chọn thật của người bán", "giữ làm nhãn (trừ số chỗ \"Khác\")"],
        ["Số km, giá là ô gõ tự do", "Gõ thiếu/thừa số 0, giá giả, giá trả trước", "quy tắc cứng + lọc giá đặt cọc"],
        ["Trang xe \"Mới\" không có ô số km", "Mọi xe mới thiếu số km", "điền 0 km"],
        ["Bảo hành ghi \"Bảo hành hãng\" ở mọi tin", "Không phân biệt được tin nào", "bỏ cột"],
        ["Một lần crawl, 750 trang mới nhất", "Ảnh chụp thị trường; giá rao, không phải giá bán", "tuổi xe tính theo 2026; nêu ở hạn chế"]]
table(s, 0.6, 1.6, 12.1, rows, [4.3, 3.9, 3.9], size=13, row_h=0.5, first_bold=False)

# ---------------------------------------------------------------- 9b. Sơ đồ quy trình chống rò rỉ dữ liệu
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Chia tập trước, học tham số chỉ từ train",
            "Câu hỏi then chốt cho mỗi bước: bước này có học con số nào từ dữ liệu không? Không học gì (quy tắc cứng, hằng số, công thức) thì làm "
            "trên toàn bộ dữ liệu ở preprocess_data. Có học (trung vị, mode, tần suất, trung vị giá) thì chia train/test trước, fit trên train rồi "
            "áp nguyên lên test ở process_data. Nếu làm ngược lại, xe trong tập test góp phần tạo ra tham số, và điểm đánh giá không còn là ước lượng "
            "cho dữ liệu mới. Tham số chỉ đi một chiều từ train sang test.")
picture(s, FIGS / "s06_so_do_quy_trinh.png", 0.3, 1.35, 6.9, 4.8)
rows = [["Bước", "Học từ dữ liệu?", "Ở đâu"],
        ["Ép kiểu, N/A → thiếu, xóa trùng", "không", "preprocess"],
        ["Giá / km vô lý (ngưỡng cứng)", "không", "preprocess"],
        ["Xe mới = 0 km, xuất xứ \"Không rõ\"", "không (hằng số)", "preprocess"],
        ["Tuổi xe, log giá", "không (công thức)", "preprocess"],
        [[("Lọc giá đặt cọc", {"bold": True})], "trung vị giá", [("process (train)", {"bold": True, "color": ORANGE})]],
        [[("Điền km, kiểu dáng, số chỗ", {"bold": True})], "trung vị, mode", [("process (train)", {"bold": True, "color": ORANGE})]],
        [[("Gộp nhãn hiếm", {"bold": True})], "tần suất", [("process (train)", {"bold": True, "color": ORANGE})]],
        ["One-hot, chuẩn hóa, tham số", "có", "model (train / CV)"]]
table(s, 7.4, 1.5, 5.4, rows, [2.55, 1.45, 1.4], size=11, row_h=0.47, first_bold=False)
box(s, 7.4, 5.85 - 0.05, 5.4, 0.45, fill=TINT_Y)
text(s, 7.5, 5.8, 5.2, 0.45, ["fit_transform(train) rồi transform(test); không bao giờ fit trên test"], size=12, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 10. Chuẩn hóa, trùng
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Chuẩn hóa giá trị giả, trích vị trí, loại tin trùng",
            "Bước đầu tiên đưa dữ liệu về đúng kiểu: chuỗi '1.095.000.000' thành số, N/A thành thiếu thật. "
            "Tin trùng hoàn toàn (cùng mọi thuộc tính và giá) bị loại vì nếu một bản ở train, một bản ở test thì điểm đánh giá cao ảo.")
cards = [("Giá trị giả", ["\"N/A\", \"Đang cập nhật\" → thiếu thật", "\"Khác\" ở số chỗ → thiếu; \"Dòng khác\" giữ làm nhãn",
                          "\"109.366\" → 109366 (bỏ dấu chấm)", "\"trước năm 1980\" → 1979"]),
         ("Vị trí và cột thừa", ["Trích tỉnh/thành từ URL: 63 tỉnh, 100% số tin",
                                 "Không tách quận/huyện: 453 giá trị, trung vị 4 tin/quận",
                                 "Bỏ no, url, veh_warranty_policy"]),
         ("Tin trùng", ["691 tin giống hệt mọi thuộc tính và giá (4,7%)", "331 trong số đó là xe mới của đại lý",
                        "Giữ một bản → còn 14.144 tin"])]
for k, (t, b) in enumerate(cards):
    x = 0.6 + k * 4.1
    box(s, x, 1.6, 3.8, 4.3, fill=TINT if k != 2 else TINT_Y)
    text(s, x + 0.2, 1.72, 3.4, 0.45, [t], size=18, color=BLUE, bold=True)
    text(s, x + 0.2, 2.3, 3.45, 3.5, b, size=14, bullets=True, space_after=8)

# ---------------------------------------------------------------- 11. Giá trị bất hợp lý
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Phát hiện giá trị bất hợp lý bằng hiểu biết thực tế",
            "Không cắt ngoại lệ theo phần trăm máy móc: ngưỡng lấy từ hiểu biết thực tế nên không học gì từ dữ liệu, làm được trên toàn bộ dữ liệu. "
            "Hình trái: giá ngoài 20 triệu - 10 tỷ bị loại (gõ thiếu / thừa số 0). Hình phải: số km xe đã dùng trên 500.000, trên 100.000 km/năm, "
            "hoặc dưới 100 km khi xe từ 2 tuổi (nhập theo nghìn km) được chuyển thành thiếu, rồi điền ở bước sau. Xe sang thật dù là ngoại lệ thống kê vẫn giữ.")
picture(s, FIGS / "p10_gia_tri_bat_hop_ly.png", 0.4, 1.4, 8.2, 3.75)
rows = [["Quy tắc", "Số tin"],
        ["Giá < 20 triệu / > 10 tỷ → loại", "64"],
        ["Giá giả 123.456.789 → loại", "14"],
        ["\"Mới\" mà > 3 năm tuổi → \"Đã sử dụng\"", "6 (sửa)"],
        ["Km > 500.000 → thiếu", "106"],
        ["Xe ≥ 2 tuổi, < 100 km → thiếu", "447"],
        ["> 100.000 km/năm → thiếu", "25"]]
table(s, 8.8, 1.5, 4.0, rows, [3.0, 1.0], size=12, row_h=0.47)
box(s, 0.6, 5.3, 12.1, 0.7, fill=TINT_Y)
text(s, 0.85, 5.33, 11.6, 0.65, [[("Giữ lại: ", {"bold": True, "color": BLUE}),
                                  ("giá 399.999.999 (định giá tâm lý), 888.888.888 (số đẹp); xe sang thật dù là ngoại lệ thống kê, xử lý bằng log giá.", {})]],
     size=14, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 13. Giá trị thiếu
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Điền thiếu: hằng số ở tiền xử lý, thống kê học từ train",
            "Xe mới thiếu số km vì trang không có ô km: điền 0 là quy tắc nghiệp vụ, làm ngay ở tiền xử lý; xuất xứ thiếu 41% nên gán nhãn hằng "
            "\"Không rõ\". Ba cột còn lại cần thống kê: số km xe cũ (trung vị theo đời xe), kiểu dáng và số chỗ (mode theo dòng xe). Các thống kê này "
            "được CarImputer học trên train ở process_data rồi áp lên test, nên file preprocess_data.csv cố ý còn giá trị thiếu.")
picture(s, FIGS / "p02_thieu_va_cach_xu_ly.png", 0.5, 1.45, 8.0, 3.45)
rows = [["Cột", "Cách điền", "Ở đâu"],
        ["Số km xe mới", "0 km", "preprocess"],
        ["Xuất xứ", "\"Không rõ\"", "preprocess"],
        ["Số km xe cũ", "trung vị đời xe + cờ", "process (train)"],
        ["Kiểu dáng", "mode dòng xe", "process (train)"],
        ["Số chỗ", "mode dòng → hãng", "process (train)"]]
table(s, 8.7, 1.55, 4.1, rows, [1.3, 1.6, 1.2], size=11, row_h=0.45)
cards = [("Thiếu có cấu trúc", "Xe mới không có ô km → điền 0; điền trung vị sẽ biến xe mới thành xe 70.000 km"),
         ("Cờ mileage_missing", "653 xe cũ ở train, 161 ở test được điền km → mô hình biết đó là ước lượng"),
         ("Không đoán bừa", "Kiểu dáng không suy ra được từ dòng xe → \"Không rõ\", không lùi về hãng")]
for k, (t, d) in enumerate(cards):
    x = 0.6 + k * 4.1
    callout(s, x, 5.0, 3.8, 1.1, t, d, size=12)

# ---------------------------------------------------------------- 13b. Chia train/test
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Chia train/test phân tầng trước khi học bất kỳ tham số nào",
            "Chia 80/20 phân tầng theo phân khúc giá để tập test có đủ xe sang: chia ngẫu nhiên thường có thể để test gần như không có xe trên 2 tỷ. "
            "Bước chia nằm đầu process_data, trước mọi bước học tham số. Không tách tập validation riêng vì chỉ có khoảng 460 xe trên 2 tỷ; "
            "CV dùng lần lượt mọi phần của train làm validation nên ổn định hơn.")
picture(s, FIGS / "m01_chia_train_test.png", 0.5, 1.5, 6.9, 3.5)
caption(s, 0.6, 5.1, 6.7, "Train 11.252, test 2.814 tin; sau lọc giá đặt cọc: 11.153 và 2.797", 14)
steps = [("Train 80%", "học mọi tham số: lọc nhãn sai, điền thiếu, gộp nhãn, mã hóa, tham số mô hình (CV / OOB)"),
         ("Test 20%", "dùng đúng một lần để báo cáo kết quả cuối"),
         ("Hai bản Random Forest", "thư viện scikit-learn làm chuẩn; bản tự viết bằng NumPy được kiểm chứng với nó")]
for k, (t, d) in enumerate(steps):
    y = 1.65 + k * 1.45
    badge(s, 7.8, y, 0.6, str(k + 1), fill=BLUE if k != 1 else ORANGE)
    text(s, 8.6, y - 0.05, 4.1, 0.4, [t], size=17, color=BLUE, bold=True)
    text(s, 8.6, y + 0.38, 4.1, 0.95, [d], size=14)

# ---------------------------------------------------------------- 12. Giá trả trước
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Giá trả trước / đặt cọc: lọc bằng trung vị của train",
            "Nhiều tin ghi số tiền trả trước thay vì giá xe. Ngưỡng tuyệt đối không bắt được: 100 triệu hợp lý với xe 20 năm tuổi nhưng vô lý với xe "
            "1 năm tuổi, nên so với chính mẫu xe đó. Trung vị giá là tham số học từ biến mục tiêu, nên chỉ tính trên train rồi áp lên cả hai tập. "
            "Vẫn xóa ở tập test vì đây là sửa nhãn sai: để nhãn sai trong test thì MAPE không còn đo đúng chất lượng mô hình. "
            "Sửa thêm một lỗi: bỏ nhóm \"Dòng khác\" khỏi bảng tham chiếu vì nhóm này trộn nhiều xe không liên quan.")
picture(s, FIGS / "s08_gia_dat_coc.png", 0.4, 1.45, 6.4, 3.6)
box(s, 7.0, 1.5, 5.7, 1.35, fill=BLUE)
text(s, 7.2, 1.55, 5.3, 0.4, ["Quy tắc (học từ train)"], size=15, color=YELLOW, bold=True)
text(s, 7.2, 1.95, 5.4, 0.9, ["Giá < 30% trung vị xe cùng hãng-dòng-năm ở train (nhóm ≥ 3 xe, bỏ \"Dòng khác\")"], size=14, color=WHITE)
rows = [["Xe", "Năm", "Giá rao", "Trung vị train"],
        ["Toyota Camry", "2026", "300", "1.360"],
        ["Ford Everest", "2026", "320", "1.310"],
        ["Toyota Fortuner (mới)", "2026", "125", "1.275"],
        ["Hyundai Santa Fe (mới)", "2025", "220", "1.079"],
        ["Mercedes E Class", "2019", "156", "899"]]
table(s, 7.0, 2.95, 5.7, rows, [2.6, 0.8, 1.0, 1.3], size=12, row_h=0.36)
cards = [("99 + 17", "tin bị loại ở train + test (cách cũ: 129 tin trên toàn bộ)"),
         ("21%", "tin test không có nhóm tham chiếu ở train → không kiểm tra được, giữ lại")]
for k, (v, d) in enumerate(cards):
    x = 0.6 + k * 6.15
    box(s, x, 5.25, 5.95, 0.8, fill=TINT)
    text(s, x + 0.15, 5.25, 1.6, 0.8, [v], size=20, color=ORANGE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.8, 5.25, 4.05, 0.8, [d], size=12, anchor=MSO_ANCHOR.MIDDLE, space_after=0)

# ---------------------------------------------------------------- 13c. Kiểm tra sau khi điền thiếu
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Kiểm tra sau khi điền: cơ cấu giữ nguyên, CV không lạc quan",
            "Kiểm tra 1: cơ cấu kiểu dáng và số chỗ ở train trước khi điền (chỉ phần đã khai) và sau khi điền gần như trùng nhau, không bị dồn về "
            "một giá trị như khi điền bằng mode toàn cục. Kiểm tra 2: train_data.csv được điền bằng thống kê của toàn bộ train, nên khi CV trên train, "
            "phần validation của mỗi fold góp phần tạo giá trị điền. Thí nghiệm đối chứng đặt CarImputer vào Pipeline để fit lại trong từng fold: "
            "MAPE chỉ chênh 0,04 điểm %, nên cách điền sẵn dùng được.")
picture(s, FIGS / "s10_truoc_sau_dien_thieu.png", 0.4, 1.45, 8.0, 3.4)
rows = [["", "Train", "Test"],
        ["Km xe cũ được điền", "653", "161"],
        ["Kiểu dáng điền theo dòng", "3.443", "884"],
        ["Kiểu dáng → \"Không rõ\"", "208", "49"],
        ["Số chỗ được điền", "2.817", "725"]]
table(s, 8.6, 1.55, 4.2, rows, [2.4, 0.9, 0.9], size=12, row_h=0.46)
box(s, 8.6, 4.05, 4.2, 1.95, fill=BLUE)
text(s, 8.8, 4.1, 3.9, 0.4, ["Đối chứng (5-fold CV, train)"], size=14, color=YELLOW, bold=True)
text(s, 8.8, 4.55, 3.9, 1.4, ["Điền sẵn: MAPE 16,81%", "Điền trong từng fold: 16,85%",
                              [("Chênh 0,04 điểm % → không đáng kể", {"bold": True, "color": YELLOW})]],
     size=13, color=WHITE, space_after=4)
caption(s, 0.6, 5.0, 7.8, "Thanh nhạt: trước khi điền (chỉ phần đã khai); thanh đậm: sau khi điền (train)", 13)

# ---------------------------------------------------------------- 14. Gộp nhãn hiếm
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Gộp nhãn hiếm theo cấp; ngưỡng chọn bằng CV trên train",
            "69 hãng và hơn 500 dòng xe, nhiều dòng chỉ 1-5 tin: one-hot nguyên vẹn sinh hàng trăm cột thưa. Gộp theo cấp: dòng xe dưới 10 tin lùi về "
            "\"<hãng> - dòng khác\" (vẫn biết là Mercedes), hãng dưới 30 tin lùi về nhóm cao cấp hoặc phổ thông theo định vị thương hiệu, không dùng giá. "
            "Đếm tần suất là học từ dữ liệu nên BrandModelGrouper chỉ đếm trên train; hãng chưa gặp ở test tự rơi về nhóm \"khác\". "
            "So 3 phương án bằng 5-fold CV: gộp theo cấp 30/10 tốt nhất ở cả hai mô hình, rõ nhất với Ridge và xe trên 2 tỷ.")
picture(s, FIGS / "s11_gop_nhan.png", 0.4, 1.4, 8.3, 3.2)
levels = [("Dòng xe ≥ 10 tin", "giữ nguyên: Toyota Vios"), ("Dòng xe hiếm", "→ \"Mercedes Benz - dòng khác\""),
          ("Hãng < 30 tin", "→ \"Hãng cao cấp khác\" (Bentley) / \"phổ thông khác\" (Wuling)")]
for k, (a, b) in enumerate(levels):
    x = 0.6 + k * 2.75
    box(s, x, 4.75, 2.55, 1.3, fill=TINT)
    badge(s, x + 0.12, 4.85, 0.42, str(k + 1), size=12)
    text(s, x + 0.62, 4.82, 1.9, 0.45, [a], size=13, color=BLUE, bold=True)
    text(s, x + 0.15, 5.3, 2.3, 0.75, [b], size=12)
callout(s, 8.95, 1.5, 3.75, 2.0, "27 hãng, 196 dòng (train)",
        ["CV MAPE (cũ → chọn): Ridge 23,8% → 19,6%; RF 17,5% → 16,6%"], dark=True, size=12)
callout(s, 8.95, 3.7, 3.75, 2.35, "Thêm vị trí có giúp không?",
        ["Thêm miền hoặc tỉnh/thành: MAPE đổi dưới 0,1 điểm % → không dùng vị trí"], size=12)

# ---------------------------------------------------------------- 15. Log giá + mất cân bằng
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Học trên log(giá); mất cân bằng theo phân khúc giá",
            "Giá lệch phải rất mạnh (skewness 5,1); log đưa về gần đối xứng (−0,26) và biến sai số tuyệt đối thành sai số tương đối, khớp với MAPE. "
            "Log là công thức tất định nên làm ở tiền xử lý. Biến mục tiêu liên tục nên không có lớp thiểu số; mất cân bằng thể hiện ở mật độ dữ liệu: "
            "xe trên 2 tỷ chỉ 3,3%, ít hơn phân khúc đông nhất gần 13 lần.")
picture(s, FIGS / "p03_gia_va_log_gia.png", 0.4, 1.45, 6.6, 3.2)
picture(s, FIGS / "p04_phan_khuc_gia.png", 7.1, 1.45, 5.7, 3.2)
callout(s, 0.6, 4.85, 5.9, 1.2, "Log giá: kiểm chứng bằng CV (phần 5)", ["MAPE 16,6% so với 20,0% khi học trên giá gốc"], dark=True, size=12)
callout(s, 6.75, 4.85, 5.95, 1.2, "Xử lý mất cân bằng",
        ["Chia phân tầng; báo MAPE theo phân khúc; thử trọng số mẫu / over / undersampling trong từng fold (phần 5)"], size=12)

# ---------------------------------------------------------------- 15b. Tổng kết dữ liệu
s = content("Phần 3 · Tiền xử lý và xử lý dữ liệu", "Từ 14.835 tin thô còn 11.153 tin train và 2.797 tin test",
            "Tiền xử lý bằng quy tắc cứng giữ 94,8% số tin. Sau khi chia tập, lọc giá đặt cọc bằng bảng của train loại thêm 99 tin train và 17 tin test. "
            "Hai file train_data.csv, test_data.csv không còn giá trị thiếu ở đặc trưng; tham số xử lý (trung vị, danh sách nhãn) lưu trong "
            "data_transformers.joblib để định giá xe mới đúng như lúc huấn luyện. File chỉ chứa dữ liệu thật, không ghi mẫu nhân bản.")
picture(s, FIGS / "p01_pheu_lam_sach.png", 0.4, 1.45, 6.9, 3.0)
rows = [["", "Train", "Test"],
        ["Sau khi chia", "11.252", "2.814"],
        ["Loại giá đặt cọc", "−99", "−17"],
        [[("Còn lại", {"bold": True})], [("11.153", {"bold": True, "color": ORANGE})], [("2.797", {"bold": True, "color": ORANGE})]],
        ["Nhóm hãng / dòng xe", "27 / 196", "27 / 193"],
        ["Còn thiếu ở đặc trưng", "0", "0"]]
table(s, 7.6, 1.55, 5.2, rows, [2.6, 1.3, 1.3], size=13, row_h=0.48, first_bold=False)
cards = [("preprocess_data.csv", "14.066 tin, còn thiếu ở 3 cột (cố ý)"),
         ("train / test_data.csv", "đã điền, gộp nhãn, có đặc trưng cho Ridge"),
         ("data_transformers.joblib", "tham số học từ train, dùng khi dự đoán")]
for k, (t, d) in enumerate(cards):
    x = 0.6 + k * 4.1
    callout(s, x, 4.75, 3.8, 1.3, t, d, dark=(k == 1), size=12)

# ---------------------------------------------------------------- 17. Thước đo
s = content("Phần 4 · Thước đo", "Thước đo đánh giá: chỉ số nào cần thấp, chỉ số nào cần cao",
            "MAPE là thước đo chính vì lệch 100 triệu là nhỏ với xe 3 tỷ nhưng rất lớn với xe 150 triệu. Không chỉ số nào tốt khi đứng một mình: "
            "luôn so với baseline luôn đoán giá trung vị. Thang MAPE là quy ước tham khảo trong dự báo (Lewis, 1982).")
rows = [["Thước đo", "Ý nghĩa", "Tốt khi", "Mức tham khảo"],
        [[("MAPE", {"bold": True, "color": ORANGE})], "sai số % trung bình (thước đo chính)", "thấp", "< 10% rất tốt · 10-20% tốt · 20-50% tạm · > 50% kém"],
        ["MedAPE", "sai số % của xe \"điển hình\" (trung vị)", "thấp", "cùng thang MAPE"],
        ["MAE", "trung bình đoán lệch bao nhiêu triệu", "thấp", "tùy mức giá"],
        ["RMSE", "như MAE nhưng phạt nặng lỗi lớn", "thấp", "RMSE ≫ MAE: có vài lỗi rất lớn"],
        ["R² (log)", "tỉ lệ biến thiên log giá được giải thích", "cao", "1 hoàn hảo · 0 = đoán trung bình"],
        ["Lệch ≤ 10% / 20%", "tỉ lệ xe đoán sát giá thật", "cao", "so với baseline"]]
table(s, 0.6, 1.6, 12.1, rows, [2.3, 4.6, 1.3, 3.9], size=13, row_h=0.52)
text(s, 0.6, 5.4, 12.1, 0.6, [[("Mốc so sánh (baseline): ", {"bold": True, "color": BLUE}),
                              ("luôn đoán giá trung vị của tập train. Luôn báo thêm MAPE theo 5 phân khúc giá.", {})]], size=14)

# ---------------------------------------------------------------- 18. Mã hóa
s = content("Phần 4 · Đặc trưng cho rừng", "One-hot hay mã số thứ tự? Kiểm chứng bằng CV",
            "One-hot sinh 258 cột 0/1; mã số thứ tự chỉ 12 cột nhưng thứ tự các nhãn là tùy ý. Với one-hot, một phép chia tách thẳng một dòng xe; "
            "với mã số, xe sang và xe rẻ xếp xen kẽ theo bảng chữ cái nên cây phải chia nhiều lần. Kết quả: one-hot tốt nhất, rõ nhất ở xe trên 2 tỷ.")
picture(s, FIGS / "m16_ma_hoa.png", 0.5, 1.45, 12.3, 3.45)
callout(s, 0.6, 5.0, 5.9, 1.1, "Chọn one-hot: MAPE 16,8%", ["mã số thứ tự 17,5-19,0%; xe > 2 tỷ 29,8% so với 30-41%"], dark=True, size=13)
callout(s, 6.8, 5.0, 5.9, 1.1, "Rừng không cần age², log, chuẩn hóa",
        ["cây chỉ so sánh ngưỡng: biến đổi đơn điệu cho cùng phép chia"], size=13)

# ---------------------------------------------------------------- 19. CART
s = content("Phần 4 · Lý thuyết", "Cây hồi quy CART: chia dữ liệu bằng câu hỏi có / không",
            "Đây là một cây thật độ sâu 2 học từ tập train. Câu hỏi đầu tiên là tuổi xe ≤ 12,5 năm, rồi đến hộp số sàn ở cả hai nhánh. "
            "Mỗi nút chọn câu hỏi làm giảm tổng bình phương sai số nhiều nhất; lá dự đoán bằng trung bình log giá của các xe trong lá.")
picture(s, FIGS / "m17_cay_minh_hoa.png", 0.3, 1.4, 10.1, 4.5)
box(s, 10.55, 1.5, 2.25, 1.9, fill=BLUE)
text(s, 10.65, 1.56, 2.1, 0.4, ["Mức giảm SSE"], size=13, color=YELLOW, bold=True)
text(s, 10.65, 1.95, 2.1, 1.4, [[("gain = S²ₗ/nₗ + S²ᵣ/nᵣ − S²/n", {"bold": True, "color": WHITE, "size": 13})],
                                [("n: số xe, S: tổng log giá", {"color": WHITE, "size": 11})]], space_after=4)
text(s, 10.55, 3.55, 2.25, 2.4, ["Cây sâu: độ chệch thấp", "Nhưng phương sai cao: đổi vài mẫu là cây đổi hẳn",
                                 "→ cần trung bình nhiều cây"], size=12, bullets=True, space_after=6)

# ---------------------------------------------------------------- 20. Rừng
s = content("Phần 4 · Lý thuyết", "Random Forest: lấy trung bình nhiều cây khác nhau",
            "Bootstrap: mỗi cây học trên mẫu lấy có hoàn lại, chứa khoảng 63% số xe khác nhau. Mỗi lần chia nút chỉ xét một phần đặc trưng ngẫu nhiên "
            "để các cây ít giống nhau. Trung bình nhiều cây ít tương quan có phương sai nhỏ hơn hẳn một cây. Khoảng 37% xe không được chọn (out-of-bag) "
            "dùng để ước lượng sai số mà không cần tập validation.")
box(s, 0.6, 2.3, 1.9, 1.15, fill=BLUE)
text(s, 0.6, 2.33, 1.9, 1.1, ["Tập train", "11.153 xe"], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
for k, lab in enumerate(["Bootstrap 1", "Bootstrap 2", "…", "Bootstrap T"]):
    y = 1.55 + k * 0.68
    arrow(s, 2.6, y + 0.13, 0.5, 0.28, fill=YELLOW)
    box(s, 3.2, y, 1.9, 0.55, fill=TINT)
    text(s, 3.2, y + 0.02, 1.9, 0.5, [lab], size=13, color=BLUE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, 5.2, y + 0.13, 0.4, 0.28, fill=YELLOW)
    box(s, 5.7, y, 1.8, 0.55, fill=TINT_Y)
    text(s, 5.7, y + 0.02, 1.8, 0.5, ["Cây " + ("T" if k == 3 else ("…" if k == 2 else str(k + 1)))], size=13, color=INK, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 7.65, 2.65, 0.5, 0.45, fill=ORANGE)
box(s, 8.3, 2.3, 1.9, 1.15, fill=ORANGE)
text(s, 8.3, 2.33, 1.9, 1.1, ["Trung bình", "T dự đoán"], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
text(s, 10.45, 1.6, 2.3, 2.6, [[("ŷ(x) = (1/T) Σ fₜ(x)", {"bold": True, "size": 15, "color": BLUE})],
                               [("trung bình T cây", {"size": 12, "color": MUTED})]], space_after=6)
ideas = [("1. Bootstrap", "mỗi cây học ~63% số xe khác nhau"), ("2. max_features", "mỗi nút chỉ xét một phần đặc trưng"),
         ("3. OOB", "~37% xe ngoài túi → ước lượng sai số")]
for k, (t, d) in enumerate(ideas):
    x = 0.6 + k * 4.1
    callout(s, x, 4.55, 3.8, 1.45, t, d, dark=(k == 2), size=13)

# ---------------------------------------------------------------- 21. RF thư viện
s = content("Phần 4 · Random Forest thư viện", "Tinh chỉnh RandomForestRegressor bằng CV trên train",
            "RandomizedSearchCV thử 12 cấu hình ngẫu nhiên, 300 cây, 3-fold CV phân tầng theo phân khúc giá. Các cấu hình tốt đều có cây sâu, lá nhỏ: "
            "dữ liệu có nhiều tổ hợp hãng-dòng-đời xe riêng biệt; việc lấy trung bình 300 cây đã kiểm soát phương sai.")
rows = [["min_samples_leaf", "max_features", "max_depth", "CV RMSE (log)"],
        [[("1", {"bold": True, "color": ORANGE})], [("0,50", {"bold": True, "color": ORANGE})], [("30", {"bold": True, "color": ORANGE})], [("0,2745", {"bold": True, "color": ORANGE})]],
        ["1", "0,33", "30", "0,2763"],
        ["2", "0,80", "30", "0,2836"],
        ["2", "0,50", "30", "0,2844"],
        ["2", "0,33", "30", "0,2876"]]
text(s, 0.6, 1.55, 7.0, 0.4, ["5 cấu hình tốt nhất (300 cây)"], size=16, color=BLUE, bold=True)
table(s, 0.6, 2.05, 7.0, rows, [2.0, 1.7, 1.5, 1.8], size=14, row_h=0.52, first_bold=False)
stats = [("0,888", "OOB R² (log) của mô hình cuối"), ("7 giây", "huấn luyện 300 cây (C/Cython, nhiều lõi)"),
         ("12 / 3-fold", "số cấu hình thử / số fold CV")]
for k, (v, l) in enumerate(stats):
    y = 1.6 + k * 1.45
    box(s, 8.0, y, 4.7, 1.25, fill=BLUE if k == 0 else TINT)
    text(s, 8.2, y + 0.08, 4.3, 0.6, [v], size=26, color=YELLOW if k == 0 else BLUE, bold=True)
    text(s, 8.2, y + 0.7, 4.3, 0.5, [l], size=13, color=WHITE if k == 0 else INK)

# ---------------------------------------------------------------- 22. RF tự cài đặt
s = content("Phần 4 · Random Forest tự cài đặt", "Tự cài đặt bằng NumPy: ba lớp, tìm ngưỡng bằng histogram",
            "Cách cơ bản phải sắp xếp từng cột ở mỗi nút, rất chậm trong Python. Kỹ thuật histogram (giống LightGBM): chia mỗi đặc trưng thành 32 bin "
            "một lần; ở mỗi nút chỉ cần np.bincount đếm số xe và tổng log giá theo bin, rồi np.cumsum cộng dồn là có gain của mọi ngưỡng cùng lúc.")
cls = [("Binner", "chia mỗi đặc trưng thành ≤ 32 bin theo phân vị (fit trên train)"),
       ("RegressionTreeScratch", "cây CART: tìm phép chia giảm SSE lớn nhất, dựng cây bằng ngăn xếp"),
       ("RandomForestScratch", "bootstrap + chọn ngẫu nhiên đặc trưng + trung bình; OOB và độ quan trọng")]
for k, (c, d) in enumerate(cls):
    y = 1.6 + k * 1.1
    badge(s, 0.6, y + 0.2, 0.5, str(k + 1), size=14)
    box(s, 1.25, y, 5.35, 0.95, fill=TINT)
    text(s, 1.4, y + 0.05, 5.1, 0.4, [c], size=15, color=BLUE, bold=True)
    text(s, 1.4, y + 0.47, 5.1, 0.45, [d], size=13)
text(s, 7.0, 1.6, 5.7, 0.4, ["Tìm ngưỡng tốt nhất ở một nút"], size=16, color=BLUE, bold=True)
hsteps = [("Đếm theo bin", "bincount: số xe, tổng log giá mỗi bin"), ("Cộng dồn", "cumsum: nhánh trái của mọi ngưỡng"),
          ("Tính gain", "công thức gain cho mọi ngưỡng cùng lúc"), ("Chọn max", "argmax trên (đặc trưng, ngưỡng)")]
for k, (t, d) in enumerate(hsteps):
    y = 2.1 + k * 0.85
    box(s, 7.0, y, 5.7, 0.72, fill=TINT_Y if k % 2 else TINT)
    text(s, 7.15, y + 0.05, 1.8, 0.62, [t], size=14, color=BLUE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.95, y + 0.05, 3.7, 0.62, [d], size=13, anchor=MSO_ANCHOR.MIDDLE)
box(s, 0.6, 5.0, 6.0, 1.0, fill=BLUE)
text(s, 0.8, 5.05, 5.6, 0.9, [[("Cột one-hot chỉ có 2 giá trị → không mất thông tin; ", {"color": WHITE}),
                               ("cột số mất rất ít độ chi tiết.", {"color": YELLOW, "bold": True})]], size=14, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 23. Kiểm chứng
s = content("Phần 4 · Kiểm chứng", "Kiểm chứng: một cây và cả rừng trùng khớp thư viện",
            "Bước 1: dựng một cây cùng tham số với DecisionTreeRegressor, so số nút và R². Bước 2: so dự đoán của cả rừng trên từng xe trong tập test. "
            "Tập test ở đây chỉ dùng để kiểm tra tính đúng của code, không để chọn tham số.")
text(s, 0.6, 1.55, 5.3, 0.4, ["Một cây (lá ≥ 5 mẫu)"], size=16, color=BLUE, bold=True)
rows = [["", "Thư viện", "Tự cài đặt"],
        ["Số nút", "3.479", "3.401"],
        ["R² (log) train", "0,922", "0,920"],
        ["R² (log) test", "0,869", "0,873"],
        ["MAPE test", "18,5%", "18,1%"]]
table(s, 0.6, 2.05, 5.3, rows, [2.1, 1.6, 1.6], size=14, row_h=0.5)
caption(s, 0.6, 4.65, 5.3, "R² train cao hơn hẳn R² test: một cây có phương sai cao", 13)
text(s, 6.2, 1.55, 6.6, 0.4, ["Cả rừng: dự đoán trên từng xe"], size=16, color=BLUE, bold=True)
picture(s, FIG / "m11_tu_cai_dat_vs_thu_vien.png", 6.1, 2.0, 6.7, 2.6)
cards = [("2,2%", "độ lệch trung vị giữa hai bản trên từng xe"), ("76%", "số xe lệch không quá 5%")]
for k, (v, d) in enumerate(cards):
    x = 6.2 + k * 3.3
    box(s, x, 4.8, 3.1, 1.2, fill=BLUE if k == 0 else TINT)
    text(s, x + 0.15, 4.85, 1.2, 1.1, [v], size=24, color=YELLOW if k == 0 else ORANGE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.35, 4.85, 1.7, 1.1, [d], size=12, color=WHITE if k == 0 else INK, anchor=MSO_ANCHOR.MIDDLE, space_after=0)

# ---------------------------------------------------------------- 24. OOB và số cây
s = content("Phần 4 · Random Forest tự cài đặt", "Từ một cây đến một rừng: sai số giảm rõ rệt",
            "Bản tự viết chọn tham số bằng OOB (6 cấu hình, mỗi cấu hình 30 cây, chỉ huấn luyện một lần), rồi huấn luyện 100 cây trong khoảng 1,5 phút. "
            "OOB RMSE giảm nhanh trong vài chục cây đầu rồi bão hòa. Trên test, rừng giảm MAPE từ khoảng 18-18,5% của một cây xuống 15,1-15,7%. "
            "OOB chỉ dùng để so sánh cấu hình: bộ mã hóa và Binner đã fit trên toàn bộ train nên OOB không hoàn toàn độc lập.")
picture(s, FIGS / "m05_rf_so_cay.png", 0.4, 1.45, 5.6, 2.95)
picture(s, FIGS / "m18_cay_vs_rung.png", 6.1, 1.5, 6.8, 2.9)
stats = [("0,354 → 0,274", "OOB RMSE: 1 cây → 100 cây"), ("0,883", "OOB R² tự cài đặt (thư viện 0,888)"),
         ("18,5% → 15,1%", "MAPE test: một cây → cả rừng")]
for k, (v, l) in enumerate(stats):
    x = 0.6 + k * 4.1
    box(s, x, 4.6, 3.8, 1.35, fill=BLUE if k == 2 else TINT)
    text(s, x + 0.2, 4.66, 3.4, 0.6, [v], size=22, color=YELLOW if k == 2 else BLUE, bold=True)
    text(s, x + 0.2, 5.3, 3.4, 0.6, [l], size=13, color=WHITE if k == 2 else INK)

# ---------------------------------------------------------------- 25. Ridge
s = content("Phần 4 · Mô hình so sánh", "Ridge Regression: mốc so sánh tuyến tính",
            "Ridge là mô hình tuyến tính tốt nhất có thể trên cùng dữ liệu: đã có age² và log(km) để biểu diễn khấu hao cong, và được tinh chỉnh bằng CV "
            "như rừng. Vì vậy phần chênh lệch giữa Ridge và Random Forest chính là đóng góp của quan hệ phi tuyến và tương tác. Ridge chỉ dùng thư viện.")
box(s, 0.6, 1.6, 6.0, 1.75, fill=TINT)
text(s, 0.8, 1.7, 5.6, 0.4, ["Hàm mất mát"], size=15, color=BLUE, bold=True)
text(s, 0.8, 2.15, 5.6, 1.1, [[("J(w) = ", {}), ("(1/2n) Σ (yᵢ − ŷᵢ)²", {"bold": True}), ("  +  ", {}),
                               ("(α/2) Σ wⱼ²", {"bold": True, "color": ORANGE})],
                              [("sai số dự đoán            khoản phạt L2", {"size": 12, "color": MUTED})]], size=20)
text(s, 0.6, 3.55, 6.0, 2.5, ["ŷ = w₀ + Σ wⱼxⱼ: tổ hợp tuyến tính, dự đoán log giá",
                              "Phạt L2 giữ hệ số ổn định khi các cột one-hot tương quan (Toyota và Toyota Vios)",
                              "α = 0,56 chọn bằng 5-fold CV trên train"], size=15, bullets=True, space_after=9)
callout(s, 7.0, 1.6, 5.7, 2.0, "Vì sao chọn Ridge làm mốc?",
        ["Cùng tập train/test, cùng thước đo, cũng được tinh chỉnh → so sánh công bằng",
         "Trả lời: phi tuyến và tương tác giúp được bao nhiêu?"], dark=True, size=13)
callout(s, 7.0, 3.8, 5.7, 2.2, "Ridge khác rừng ở đặc trưng số",
        ["Cần age², log(1+km), log(1+km/năm) và StandardScaler",
         "Vì chỉ cộng hiệu ứng tuyến tính, trong khi xe mất giá nhanh những năm đầu rồi chậm lại"], size=13)

# ---------------------------------------------------------------- 26. Kết quả tổng thể
s = content("Phần 5 · Kết quả", "Random Forest: MAPE 15,1%, 79% xe lệch không quá 20%",
            "Đánh giá trên 2.797 tin test, dùng một lần; tập test không góp phần vào bất kỳ tham số nào. Random Forest vượt xa baseline và tốt hơn Ridge "
            "3,3 điểm %. MedAPE 8,8% thấp hơn MAPE nhiều: xe điển hình chỉ lệch khoảng 9%, một nhóm nhỏ xe sai nặng kéo trung bình lên. "
            "RMSE gấp 3 lần MAE cho thấy có vài lỗi rất lớn, chủ yếu xe sang.")
big = [("88,0%", "Baseline", GRAY), ("18,5%", "Ridge (so sánh)", BLUE), ("15,1%", "Random Forest", ORANGE)]
for k, (v, l, col) in enumerate(big):
    x = 0.6 + k * 4.1
    box(s, x, 1.55, 3.8, 1.35, fill=col)
    text(s, x, 1.58, 3.8, 0.8, [v], size=36, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, 2.35, 3.8, 0.45, [l + " · MAPE"], size=14, color=WHITE, align=PP_ALIGN.CENTER)
rows = [["Mô hình", "MAPE", "MedAPE", "MAE (tr)", "RMSE (tr)", "R² (log)", "≤ 20%", "Huấn luyện"],
        ["Baseline (trung vị)", "88,0%", "40,4%", "323", "642", "0,00", "26%", "–"],
        ["Ridge – thư viện", "18,5%", "10,5%", "108", "313", "0,878", "75%", "0,04 giây"],
        [[("Random Forest – thư viện", {"bold": True, "color": ORANGE})], [("15,1%", {"bold": True})], "8,8%", "90", "270", "0,910", "79%", "7 giây"],
        ["Random Forest – tự cài đặt", "15,7%", "9,5%", "95", "289", "0,904", "78%", "86 giây"]]
table(s, 0.6, 3.15, 12.1, rows, [3.5, 1.1, 1.2, 1.2, 1.3, 1.2, 1.1, 1.5], size=14, row_h=0.5)
caption(s, 0.6, 5.7, 12.1, "Theo thang tham khảo, MAPE 10-20% là mức tốt. Bản tự cài đặt kém bản thư viện 0,6 điểm % và chậm hơn khoảng 12 lần.", 13)

# ---------------------------------------------------------------- 27. Phân khúc
s = content("Phần 5 · Kết quả", "Mô hình tốt ở phân khúc giữa, kém ở hai đầu thang giá",
            "Phân khúc 200 triệu - 1 tỷ chiếm khoảng 75% thị trường, sai số khoảng 11-12%. Xe dưới 200 triệu: chênh vài chục triệu đã là phần trăm lớn, "
            "và còn sót tin ghi giá trả trước. Xe trên 2 tỷ: ít dữ liệu và thiếu thông tin phiên bản, trang bị.")
picture(s, FIGS / "m07_mape_phan_khuc.png", 0.5, 1.45, 12.3, 3.65)
cards = [("~11%", "sai số ở phân khúc 200 triệu - 1 tỷ (75% thị trường)"),
         ("32%", "xe dưới 200 triệu: giá thấp, tình trạng khác nhau, giá trả trước sót"),
         ("21%", "xe trên 2 tỷ: chỉ 3% dữ liệu, thiếu phiên bản / trang bị")]
for k, (v, d) in enumerate(cards):
    x = 0.6 + k * 4.1
    box(s, x, 5.2, 3.8, 0.85, fill=TINT)
    text(s, x + 0.15, 5.25, 1.2, 0.75, [v], size=22, color=ORANGE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.35, 5.25, 2.35, 0.75, [d], size=12, anchor=MSO_ANCHOR.MIDDLE, space_after=0)

# ---------------------------------------------------------------- 28. Dự đoán so với thực tế
s = content("Phần 5 · Kết quả", "Dự đoán so với giá thật: RF bám sát đường chéo",
            "Mỗi điểm là một xe test; trục thang log nên khoảng cách tới đường chéo là sai số phần trăm. Vùng xám là ±20%. "
            "Phần dư đối xứng quanh 0 nên mô hình không đoán lệch hệ thống; Random Forest có phân bố hẹp hơn Ridge.")
picture(s, FIGS / "m08_du_doan_vs_thuc_te.png", 0.4, 1.4, 8.7, 3.85)
text(s, 9.3, 1.6, 3.4, 4.0, ["R² (log): Ridge 0,878, RF 0,910", "Lệch ≤ 20%: Ridge 75%, RF 79%",
                             "Không đoán lệch hệ thống (phần dư quanh 0)", "Điểm xa đường chéo nhất: xe rất rẻ và xe sang"],
     size=14, bullets=True, space_after=10)
caption(s, 0.6, 5.3, 12.1, "Giá dự đoán = exp(dự đoán log giá), tức ước lượng trung vị của giá; phù hợp với MAPE.", 13)

# ---------------------------------------------------------------- 29. Bootstrap
s = content("Phần 5 · Kết quả", "Chênh lệch có ý nghĩa thống kê không? Bootstrap ghép cặp",
            "Khoảng tin cậy 95% cho biết MAPE có thể dao động trong khoảng nào nếu đánh giá trên tập test khác. Cách tính: bốc lại 2.797 xe test có hoàn lại, "
            "tính MAPE, lặp 1.000 lần, bỏ 2,5% mỗi đầu. So hai mô hình: tính hiệu MAPE trên cùng các xe trong mỗi lần bốc. Khoảng của hiệu không chứa 0 "
            "nghĩa là chênh lệch có ý nghĩa.")
picture(s, FIGS / "m10_khoang_tin_cay.png", 0.4, 1.55, 7.5, 3.4)
caption(s, 0.6, 5.0, 7.2, "Bốc lại tập test 1.000 lần → khoảng chứa 95% giá trị MAPE ở giữa", 13)
res = [("RF tốt hơn Ridge", "3,3 điểm %", "KTC 95%: 2,4 → 4,4", "không chứa 0 → có ý nghĩa", BLUE),
       ("RF tự cài đặt − thư viện", "0,6 điểm %", "KTC 95%: 0,34 → 0,86", "có ý nghĩa nhưng nhỏ", TINT)]
for k, (t, v, ci, verdict, col) in enumerate(res):
    y = 1.6 + k * 1.75
    dark = col == BLUE
    box(s, 8.1, y, 4.6, 1.55, fill=col)
    text(s, 8.3, y + 0.08, 4.2, 0.4, [t], size=15, color=WHITE if dark else BLUE, bold=True)
    text(s, 8.3, y + 0.55, 2.0, 0.6, [v], size=22, color=YELLOW if dark else ORANGE, bold=True)
    text(s, 10.3, y + 0.5, 2.35, 0.95, [ci, verdict], size=12, color=WHITE if dark else INK, space_after=2)
box(s, 8.1, 5.1, 4.6, 0.9, fill=TINT_Y)
text(s, 8.25, 5.12, 4.3, 0.85, ["Quy tắc: khoảng của hiệu không chứa 0 thì chênh lệch là thật"], size=13, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 30. Độ quan trọng
s = content("Phần 5 · Diễn giải", "Đặc trưng nào quyết định giá? Hai cách đo",
            "Impurity cộng mức giảm sai số khi học, nhưng bị lệch khi đặc trưng tương quan. Permutation xáo trộn từng đặc trưng trên tập test rồi đo MAPE "
            "tăng bao nhiêu. Hộp số tụt từ hạng 2 xuống hạng 4 vì trùng thông tin với tuổi xe: xe số sàn thường là xe cũ.")
picture(s, FIGS / "m12_do_quan_trong.png", 0.6, 1.4, 12.1, 3.6)
cards = [("Tuổi xe đứng đầu", "xáo trộn làm MAPE tăng 45 điểm %"),
         ("Hộp số: hạng 2 → 4", "trùng thông tin với tuổi xe"),
         ("Hãng: hạng 3 → 2", "giá thương hiệu khó thay thế")]
for k, (t, d) in enumerate(cards):
    x = 0.6 + k * 4.1
    callout(s, x, 5.05, 3.8, 1.0, t, d, dark=(k == 0), size=12)

# ---------------------------------------------------------------- 31. Khấu hao
s = content("Phần 5 · Diễn giải", "Random Forest học được mỗi dòng xe khấu hao khác nhau",
            "Cố định mọi thuộc tính, chỉ thay đổi tuổi xe với 15.000 km mỗi năm. Ridge buộc mọi dòng xe mất cùng một tỉ lệ % mỗi năm. "
            "Random Forest học được Mercedes C Class mất giá rất nhanh 3-4 năm đầu: đó là tương tác giữa dòng xe và tuổi xe.")
picture(s, FIGS / "m13_duong_khau_hao.png", 0.5, 1.45, 12.3, 3.8)
text(s, 0.6, 5.35, 12.1, 0.75, [[("Vì sao Random Forest chính xác hơn: ", {"bold": True, "color": BLUE}),
                                 ("giá phụ thuộc phi tuyến và có tương tác giữa các đặc trưng. Mức chênh 3,3 điểm % chính là phần mô hình tuyến tính bỏ lỡ.", {})]],
     size=15)

# ---------------------------------------------------------------- 32. Log + mất cân bằng
s = content("Phần 5 · Thí nghiệm", "Log giá và xử lý mất cân bằng: kiểm chứng bằng CV trên train",
            "Log giá giảm MAPE từ 20,0% xuống 16,6%. Trọng số mẫu và oversampling không cải thiện MAPE tổng nhưng giảm sai số xe trên 2 tỷ khoảng 3 điểm %. "
            "Undersampling bỏ 22% dữ liệu và làm kết quả kém đi. Cân bằng chỉ làm trong phần train của từng fold; đánh giá bằng CV vì oversampling làm OOB sai.")
picture(s, FIGS / "m14_log_va_mat_can_bang.png", 0.5, 1.45, 12.3, 3.6)
cards = [("Log giá", "20,0% → 16,6% MAPE tổng"), ("Oversampling", "xe > 2 tỷ: 28,4% → 25,0%"), ("Undersampling", "mất 22% dữ liệu, MAPE 17,2%")]
for k, (t, d) in enumerate(cards):
    x = 0.6 + k * 4.1
    box(s, x, 5.15, 3.8, 0.9, fill=TINT if k != 0 else BLUE)
    text(s, x + 0.2, 5.2, 3.5, 0.4, [t], size=15, color=YELLOW if k == 0 else BLUE, bold=True)
    text(s, x + 0.2, 5.58, 3.5, 0.4, [d], size=13, color=WHITE if k == 0 else INK)

# ---------------------------------------------------------------- 33. Tin đăng lại, nhãn sai
s = content("Phần 5 · Hạn chế", "Hai nguồn sai lệch: tin đăng lại và nhãn sai còn sót",
            "Tin đăng lại làm điểm hơi lạc quan: chia theo nhóm (mọi bản đăng của một xe cùng một phía) làm MAPE CV tăng 0,3 điểm %. "
            "Ngược lại, nhiều xe bị đoán sai nhất thực ra là tin ghi giá trả trước còn sót: mô hình đoán hợp lý hơn giá rao.")
text(s, 0.6, 1.55, 5.6, 0.4, ["Chia ngẫu nhiên so với chia theo nhóm (CV)"], size=15, color=BLUE, bold=True)
rows = [["Cách chia", "MAPE", "< 200 tr", "> 2 tỷ"],
        ["Ngẫu nhiên", "16,6%", "39,7%", "28,4%"],
        ["Theo nhóm", "17,0%", "41,1%", "29,6%"]]
table(s, 0.6, 2.0, 5.6, rows, [2.0, 1.2, 1.2, 1.2], size=14, row_h=0.5)
callout(s, 0.6, 3.75, 5.6, 2.25, "Ước lượng thận trọng: khoảng 16%",
        ["14% tin test có bản sao trong train: MAPE 10,1% so với 15,9% ở phần còn lại",
         "Mức lạc quan có nhưng nhỏ, không đổi kết luận"], dark=True, size=13)
text(s, 6.5, 1.55, 6.2, 0.4, ["Xe bị đoán sai nhiều nhất (triệu đồng)"], size=15, color=BLUE, bold=True)
rows = [["Xe", "Tình trạng", "Giá rao", "Dự đoán"],
        ["Hyundai Accent 2026", "Mới", "90", "465"],
        ["Geely (dòng khác) 2026", "Mới", "114", "604"],
        ["GAC M8 2025", "Mới", "168", "621"],
        ["MINI Cooper 2025", "Đã dùng", "199", "751"],
        ["VinFast (dòng khác) 2025", "Đã dùng", "200", "590"]]
table(s, 6.5, 2.0, 6.2, rows, [2.8, 1.3, 1.0, 1.1], size=13, row_h=0.47)
caption(s, 6.5, 4.95, 6.2, "Xe đời mới rao 90-200 triệu: gần như chắc chắn là giá trả trước lọt lưới (nhóm < 3 xe ở train hoặc \"Dòng khác\").", 12)

# ---------------------------------------------------------------- 33b. Trước / sau khi sửa quy trình
s = content("Phần 5 · Quy trình", "Trước / sau khi sửa quy trình: kết luận không đổi",
            "Bản cũ điền thiếu, gộp nhãn và lọc giá đặt cọc trên toàn bộ dữ liệu rồi mới chia tập. Bản mới chia tập trước, mọi tham số học từ train. "
            "Phần chênh do thống kê chỉ học từ train đo bằng thí nghiệm đối chứng: 0,04 điểm %. Phần còn lại do tập test được chia lại khác đi "
            "(bỏ lọc ở tiền xử lý làm số dòng trước khi chia thay đổi), nằm trong khoảng tin cậy bootstrap khoảng ±1 điểm %. "
            "Mọi kết luận giữ nguyên: thứ hạng mô hình, log giá, one-hot, cân bằng, tuổi xe quan trọng nhất.")
picture(s, FIGS / "s07_truoc_sau.png", 0.4, 1.4, 8.3, 3.4)
rows = [["MAPE test", "Trước", "Sau"],
        ["Ridge", "18,5%", "18,5%"],
        [[("RF thư viện", {"bold": True})], "15,5%", [("15,1%", {"bold": True, "color": ORANGE})]],
        ["RF tự cài đặt", "15,7%", "15,7%"],
        ["Số tin test", "2.788", "2.797"]]
table(s, 8.9, 1.5, 3.85, rows, [1.75, 1.05, 1.05], size=13, row_h=0.46)
box(s, 8.9, 4.0, 3.85, 0.8, fill=BLUE)
text(s, 9.05, 4.0, 3.6, 0.8, [[("Đối chứng: chỉ 0,04 điểm %", {"bold": True, "color": YELLOW})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
cards = [("Trước", "Tham số tiền xử lý tính trên toàn bộ dữ liệu: tập test góp phần tạo giá trị điền, nhãn gộp, bảng lọc"),
         ("Sau", "Tham số chỉ học từ train; test chỉ được transform → điểm là ước lượng trung thực cho xe mới")]
for k, (t, d) in enumerate(cards):
    x = 0.6 + k * 6.15
    callout(s, x, 5.0, 5.95, 1.05, t, d, dark=(k == 1), size=12)

# ---------------------------------------------------------------- 34. Lưu mô hình, demo
s = content("Phần 5 · Triển khai", "Lưu mô hình và dự đoán từ thông tin thô của xe",
            "Mô hình được lưu trong thư mục model/. Hàm predict_price nhận thông tin thô (hãng, dòng, năm, km...), tự gộp nhãn, điền km thiếu "
            "bằng đúng tham số đã học từ train (lưu trong metadata.json) và dựng đặc trưng như lúc huấn luyện. "
            "Nạp lại cho dự đoán trùng notebook đến 10⁻¹¹ triệu, kể cả rừng tự cài đặt.")
rows = [["File", "Nội dung"],
        ["random_forest_sklearn.joblib", "rừng thư viện (300 cây)"],
        ["random_forest_scratch.npz", "các cây và ngưỡng bin của rừng tự cài đặt"],
        ["preprocessor_rf.joblib", "bộ mã hóa one-hot (fit trên train)"],
        ["metadata.json", "nhãn đủ ngưỡng, trung vị km (học từ train), metric"],
        ["data_transformers.joblib", "CarImputer, BrandModelGrouper đã fit trên train"],
        ["price_model.py", "load_model(), predict_price()"]]
table(s, 0.6, 1.6, 6.6, rows, [3.1, 3.5], size=12, row_h=0.47)
text(s, 7.5, 1.55, 5.2, 0.4, ["Ví dụ dự đoán (triệu đồng)"], size=15, color=BLUE, bold=True)
rows = [["Xe", "RF (TV)", "RF (TCĐ)", "Ridge"],
        ["Toyota Vios 2018, 80.000 km", "299", "300", "283"],
        ["Ford Ranger 2021, không rõ km", "518", "563", "600"]]
table(s, 7.5, 2.0, 5.2, rows, [2.6, 0.9, 0.9, 0.8], size=13, row_h=0.6)
box(s, 7.5, 4.0, 5.2, 1.9, fill=TINT_Y)
text(s, 7.7, 4.08, 4.9, 1.8, ["Xe đã dùng không rõ km → điền trung vị km theo năm sản xuất, đánh dấu mileage_missing",
                              "Dòng xe chưa gặp → lùi về \"<hãng> - dòng khác\""], size=13, bullets=True, space_after=6)

# ---------------------------------------------------------------- 35. Kết luận
s = content("Kết luận", "Kết luận, hạn chế và hướng phát triển",
            "Tóm tắt kết quả, nêu rõ hạn chế và hướng phát triển. Cảm ơn thầy cô đã lắng nghe, nhóm em sẵn sàng trả lời câu hỏi.")
pts = [("15,1%", "MAPE của Random Forest; 79% xe lệch ≤ 20%"), ("+3,3 điểm %", "tốt hơn Ridge, có ý nghĩa thống kê"),
       ("15,7%", "rừng tự cài đặt, kiểm chứng ở mức cây và rừng")]
for k, (v, d) in enumerate(pts):
    x = 0.6 + k * 4.1
    box(s, x, 1.55, 3.8, 1.35, fill=BLUE if k == 0 else TINT)
    text(s, x + 0.2, 1.6, 3.4, 0.65, [v], size=26, color=YELLOW if k == 0 else ORANGE, bold=True)
    text(s, x + 0.2, 2.25, 3.45, 0.65, [d], size=13, color=WHITE if k == 0 else INK)
box(s, 0.6, 3.05, 5.95, 2.35, fill=TINT)
text(s, 0.85, 3.12, 5.5, 0.4, ["Hạn chế"], size=16, color=BLUE, bold=True)
text(s, 0.85, 3.55, 5.5, 1.8, ["Giá rao bán, không phải giá giao dịch",
                               "Chưa lấy phiên bản, số đời chủ, đăng kiểm",
                               "Còn sót giá trả trước ở dòng xe ít tin; có tin đăng lại",
                               "Ảnh chụp thị trường 10/2026, xe ≤ 10 tỷ"], size=14, bullets=True, space_after=5)
box(s, 6.75, 3.05, 5.95, 2.35, fill=TINT_Y)
text(s, 7.0, 3.12, 5.5, 0.4, ["Hướng phát triển"], size=16, color=BLUE, bold=True)
text(s, 7.0, 3.55, 5.5, 1.8, ["Crawl tiêu đề, mô tả: phiên bản xe, lọc giá trả trước",
                              "Chia train/test theo nhóm làm đánh giá chính",
                              "Gradient Boosting dùng lại lớp cây tự cài đặt",
                              "Giao diện web nhập thông tin xe, nhận giá"], size=14, bullets=True, space_after=5)
text(s, 0.6, 5.5, 12.1, 0.45, [[("Cảm ơn thầy cô đã lắng nghe!", {"bold": True, "color": ORANGE})]], size=18, align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------- xóa slide mẫu, sắp xếp
sldIdLst = prs.slides._sldIdLst
keep = {id(sl._element) for sl in order}
id_of = {}
for sldId in list(sldIdLst):
    part = prs.part.related_part(sldId.rId)
    if id(part._element) not in keep:
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)
    else:
        id_of[id(part._element)] = sldId
for sl in order:
    el = id_of[id(sl._element)]
    sldIdLst.remove(el)
    sldIdLst.append(el)

prs.save(OUT)
print("slides:", len(prs.slides))
