"""Dựng báo cáo Word cho đề tài dự đoán giá ô tô.

Chạy: python build_report.py   (thư mục làm việc bất kỳ)
Số liệu lấy từ output đã lưu của src/preprocess_data.ipynb và src/model.ipynb, không chạy lại notebook.
Sau khi tạo file, mở bằng Word và cập nhật mục lục (chuột phải > Update Field) nếu mục lục trống.
"""
import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "src" / "figures"
LOGO = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "logo_dut.jpeg"
OUT = ROOT / "report" / "Bao_cao_Du_doan_gia_o_to.docx"
XSL = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

BLUE = RGBColor(0x07, 0x52, 0x89)
FONT = "Times New Roman"

doc = Document()

# ------------------------------------------------------------------ styles
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(13)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = st.paragraph_format
pf.line_spacing = 1.3
pf.space_after = Pt(6)
pf.space_before = Pt(0)

for name, size, align in [("Heading 1", 15, WD_ALIGN_PARAGRAPH.CENTER),
                          ("Heading 2", 13.5, WD_ALIGN_PARAGRAPH.LEFT),
                          ("Heading 3", 13, WD_ALIGN_PARAGRAPH.LEFT)]:
    h = doc.styles[name]
    h.font.name = FONT
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.italic = name == "Heading 3"
    h.font.color.rgb = BLUE if name != "Heading 3" else RGBColor(0, 0, 0)
    rpr = h.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    h.paragraph_format.alignment = align
    h.paragraph_format.space_before = Pt(18 if name == "Heading 1" else 12)
    h.paragraph_format.space_after = Pt(12 if name == "Heading 1" else 6)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.line_spacing = 1.15

from docx.enum.style import WD_STYLE_TYPE
for name in ("Chú thích hình", "Chú thích bảng"):
    s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = st
    s.font.size = Pt(12)
    s.font.italic = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_before = Pt(4)
    s.paragraph_format.space_after = Pt(10)
    if name == "Chú thích bảng":
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(8)
        s.paragraph_format.space_after = Pt(4)

# ------------------------------------------------------------------ page
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_runs(p, text, size=None, color=None, bold=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            add_runs(p, part[2:-2], size, color, True)   # cho phép `code` nằm trong **đậm**
            continue
        if part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 13) - 1.5)
        elif part.startswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if size and not part.startswith("`"):
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        if bold:
            r.bold = True
    return p


def para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, **kw):
    p = doc.add_paragraph()
    p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(1)
    return add_runs(p, text, **kw)


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    return add_runs(p, text)


def numbered(items):
    for k, t in enumerate(items, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        add_runs(p, f"{k}. " + t)


def h1(t):
    doc.add_heading(t, level=1)


def h2(t):
    doc.add_heading(t, level=2)


def h3(t):
    doc.add_heading(t, level=3)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)


TAB_N = {}
FIG_N = {}
CHAPTER = [0]


def caption_table(text):
    c = CHAPTER[0]
    TAB_N[c] = TAB_N.get(c, 0) + 1
    p = doc.add_paragraph(style="Chú thích bảng")
    r = p.add_run(f"Bảng {c}.{TAB_N[c]}. "); r.bold = True; r.italic = False
    p.add_run(text)


def table(headers, rows, widths, caption=None, bold_rows=(), align=None, font=11.5):
    if caption:
        caption_table(caption)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_margins(t)
    align = align or ["l"] + ["c"] * (len(headers) - 1)
    amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
    for i, row in enumerate([headers] + rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else amap[align[j]]
            add_runs(p, str(val), size=font, bold=(i == 0 or (i - 1) in bold_rows) or None,
                     color=RGBColor(255, 255, 255) if i == 0 else None)
            if i == 0:
                shade(cell, "075289")
            elif (i - 1) in bold_rows:
                shade(cell, "FFF6CC")
            elif i % 2 == 0:
                shade(cell, "EAF2FA")
    # lặp tiêu đề khi sang trang
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for col, w in zip(t.columns, widths):
        col.width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(name, width_cm, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(FIG / name), width=Cm(width_cm))
    c = CHAPTER[0]
    FIG_N[c] = FIG_N.get(c, 0) + 1
    cp = doc.add_paragraph(style="Chú thích hình")
    r = cp.add_run(f"Hình {c}.{FIG_N[c]}. "); r.bold = True; r.italic = False
    cp.add_run(text)


_xslt = etree.XSLT(etree.parse(str(XSL)))


def equation(latex):
    mml = etree.fromstring(latex_to_mathml(latex).encode())
    omml = _xslt(mml).getroot()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p._p.append(parse_xml(etree.tostring(omml)))


def field(p, instr, placeholder="Nhấn chuột phải, chọn Update Field để cập nhật."):
    def fc(t):
        r = OxmlElement("w:r"); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r.append(e); return r
    p._p.append(fc("begin"))
    r = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r.append(it); p._p.append(r)
    p._p.append(fc("separate"))
    r = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = placeholder; r.append(tt); p._p.append(r)
    p._p.append(fc("end"))


def chapter(n, title):
    CHAPTER[0] = n
    page_break()
    h1(f"CHƯƠNG {n}. {title}")


def front_heading(t):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(t); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = BLUE


# ================================================================== BÌA
def cover_line(text, size, bold=False, after=0, color=None, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


cover_line("ĐẠI HỌC ĐÀ NẴNG", 14, True)
cover_line("TRƯỜNG ĐẠI HỌC BÁCH KHOA", 14, True)
cover_line("KHOA <TÊN KHOA>", 14, True, after=4)
cover_line("─────── ✦ ───────", 12, after=24)
if LOGO and LOGO.exists():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    p.add_run().add_picture(str(LOGO), width=Cm(3.6))
cover_line("BÁO CÁO ĐỀ TÀI", 16, True, after=4)
cover_line("MÔN KHOA HỌC DỮ LIỆU", 16, True, after=30)
cover_line("DỰ ĐOÁN GIÁ Ô TÔ TRÊN CHỢ TỐT", 22, True, color=BLUE, after=4)
cover_line("BẰNG MÔ HÌNH RANDOM FOREST", 18, True, color=BLUE, after=60)

info = doc.add_table(rows=4, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate([("Giảng viên hướng dẫn:", "<Tên giảng viên>"),
                            ("Sinh viên thực hiện:", "<Họ tên SV 1> – <MSSV>"),
                            ("", "<Họ tên SV 2> – <MSSV>"),
                            ("Lớp học phần:", "<Mã lớp>")]):
    for j, (val, w) in enumerate([(k, 6), (v, 7.5)]):
        c = info.cell(i, j); c.width = Cm(w)
        pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(2)
        r = pp.add_run(val); r.font.size = Pt(14); r.bold = j == 0
for _ in range(5):
    doc.add_paragraph()
cover_line("Đà Nẵng, tháng 10 năm 2026", 14, italic=True)

# khung trang bìa
sectPr = doc.sections[0]._sectPr
pgb = parse_xml(
    '<w:pgBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:offsetFrom="page">'
    + "".join(f'<w:{s} w:val="thinThickSmallGap" w:sz="24" w:space="24" w:color="075289"/>'
              for s in ("top", "left", "bottom", "right")) + "</w:pgBorders>")
sectPr.append(pgb)

# sang section mới (không khung, có số trang)
new = doc.add_section(WD_SECTION.NEW_PAGE)
new_pr = new._sectPr
for e in new_pr.findall(qn("w:pgBorders")):
    new_pr.remove(e)
new.footer.is_linked_to_previous = False
fp = new.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
field(fp, "PAGE", "1")
pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); new_pr.append(pg)
doc.sections[0].footer.is_linked_to_previous = False

# ================================================================== MỤC LỤC
front_heading("MỤC LỤC")
field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
page_break()
front_heading("DANH MỤC HÌNH")
field(doc.add_paragraph(), 'TOC \\h \\z \\t "Chú thích hình,1"')
page_break()
front_heading("DANH MỤC BẢNG")
field(doc.add_paragraph(), 'TOC \\h \\z \\t "Chú thích bảng,1"')
page_break()
front_heading("DANH MỤC TỪ VIẾT TẮT VÀ THUẬT NGỮ")
CHAPTER[0] = 0
table(["Từ viết tắt", "Tiếng Anh", "Ý nghĩa"], [
    ["MAPE", "Mean Absolute Percentage Error", "Sai số phần trăm tuyệt đối trung bình (thước đo chính)"],
    ["MedAPE", "Median Absolute Percentage Error", "Trung vị sai số phần trăm tuyệt đối"],
    ["MAE", "Mean Absolute Error", "Sai số tuyệt đối trung bình (triệu đồng)"],
    ["RMSE", "Root Mean Squared Error", "Căn bậc hai sai số bình phương trung bình"],
    ["R²", "Coefficient of determination", "Hệ số xác định"],
    ["CV", "Cross-validation", "Kiểm định chéo (K-fold)"],
    ["OOB", "Out-of-bag", "Các mẫu không được chọn vào mẫu bootstrap của một cây"],
    ["CART", "Classification and Regression Tree", "Cây phân loại và hồi quy"],
    ["RF", "Random Forest", "Rừng ngẫu nhiên"],
    ["SSE", "Sum of Squared Errors", "Tổng bình phương sai số"],
    ["KTC", "Confidence interval", "Khoảng tin cậy"],
    ["EDA", "Exploratory Data Analysis", "Phân tích khám phá dữ liệu"],
], [2.6, 5.8, 7.6], align=["l", "l", "l"])

# ================================================================== CHƯƠNG 1
chapter(1, "GIỚI THIỆU ĐỀ TÀI")
h2("1.1. Bối cảnh")
para("Thị trường ô tô cũ ở Việt Nam ngày càng sôi động, và phần lớn giao dịch bắt đầu từ các trang rao vặt trực tuyến như Chợ Tốt (`xe.chotot.com`). Người bán thường tự đặt giá theo cảm tính hoặc tham khảo vài tin tương tự, còn người mua khó biết một mức giá rao là hợp lý hay không. Một mô hình ước lượng giá từ các thông tin cơ bản của xe (hãng, dòng xe, năm sản xuất, số km, hộp số, nhiên liệu, kiểu dáng...) giúp cả hai phía có một mốc tham chiếu khách quan.")
para("Đề tài được thực hiện theo yêu cầu của môn Khoa học dữ liệu: xây dựng **một mô hình học máy từ đầu đến cuối**, từ thu thập dữ liệu, tiền xử lý, huấn luyện đến đánh giá bằng các thước đo phù hợp, có áp dụng **kỹ thuật xử lý dữ liệu mất cân bằng**. Mô hình phải được **phân tích kỹ** và **tự cài đặt** thuật toán, không chỉ gọi thư viện. Đề tài chọn **Random Forest** làm mô hình chính.")

h2("1.2. Phát biểu bài toán")
para("Đây là bài toán **hồi quy có giám sát**:")
bullet("**Đầu vào:** các thuộc tính của một tin rao bán ô tô: hãng, dòng xe, xuất xứ, tình trạng (mới / đã sử dụng), hộp số, nhiên liệu, kiểu dáng, số chỗ, năm sản xuất, số km đã đi.")
bullet("**Đầu ra:** giá rao bán dự đoán (triệu đồng).")
bullet("**Dữ liệu:** 14.835 tin rao bán ô tô thu thập từ Chợ Tốt ngày 25/10/2026; sau tiền xử lý còn 13.937 tin.")
bullet("**Thước đo chính:** MAPE (sai số phần trăm trung bình), báo cáo kèm theo từng phân khúc giá.")

h2("1.3. Mục tiêu và phạm vi")
numbered([
    "Thu thập dữ liệu tin rao bán ô tô từ Chợ Tốt bằng web scraping.",
    "Làm sạch dữ liệu: xử lý giá trị thiếu, giá trị bất hợp lý, giá giả, giá trả trước, tin trùng; gộp nhãn hiếm; tạo đặc trưng.",
    "Xây dựng mô hình chính **Random Forest** với hai bản: dùng thư viện `scikit-learn` và **tự cài đặt bằng NumPy**, rồi kiểm chứng hai bản cho kết quả tương đương. **Ridge Regression** (thư viện) chỉ dùng làm mô hình so sánh.",
    "Đánh giá bằng nhiều thước đo, kiểm định ý nghĩa thống kê của chênh lệch giữa các mô hình.",
    "Thử nghiệm các kỹ thuật xử lý mất cân bằng (trọng số mẫu, oversampling, undersampling) và đánh giá tác động.",
    "Diễn giải mô hình: yếu tố nào quyết định giá xe, xe mất giá theo tuổi như thế nào.",
])
para("**Phạm vi:** giá trong dữ liệu là **giá rao bán**, không phải giá giao dịch thực tế. Dữ liệu là ảnh chụp thị trường tại một thời điểm, không xét biến động giá theo thời gian.")

h2("1.4. Quy trình thực hiện")
para("Đề tài đi theo quy trình chuẩn của một dự án khoa học dữ liệu, mỗi bước tương ứng với một notebook:")
table(["Bước", "Nội dung", "File", "Kết quả"], [
    ["1. Thu thập", "Crawl 750 trang danh sách tin, đọc thuộc tính từng tin", "`crawl_data.ipynb`", "`raw_data_oto.csv` (14.835 dòng, 14 cột)"],
    ["2. Khám phá", "Biểu đồ phân bố, tương quan, giá trị thiếu", "`visualize_raw_data.ipynb`", "Biểu đồ trong `figures/`"],
    ["3. Tiền xử lý", "Làm sạch, điền thiếu, gộp nhãn hiếm, tạo đặc trưng", "`preprocess_data.ipynb`", "`clean_data_oto.csv` (13.937 dòng, 18 cột)"],
    ["4. Mô hình", "Random Forest (thư viện + tự cài đặt), Ridge để so sánh, đánh giá, thí nghiệm", "`model.ipynb`", "Mô hình đã lưu trong `model/`"],
    ["5. Triển khai", "Nạp mô hình, dự đoán từ thông tin thô của xe", "`model/price_model.py`", "Hàm `predict_price()`"],
], [2.4, 5.6, 4.0, 4.0], caption="Các bước thực hiện đề tài", align=["l", "l", "l", "l"])

h2("1.5. Công cụ sử dụng")
bullet("**Ngôn ngữ:** Python 3, môi trường Jupyter Notebook.")
bullet("**Thu thập dữ liệu:** `requests`, `BeautifulSoup`.")
bullet("**Xử lý và trực quan hóa:** `pandas` 3.0, `NumPy`, `matplotlib`.")
bullet("**Học máy:** `scikit-learn` 1.9 (bản thư viện, dùng làm chuẩn đối chiếu); `NumPy` (bản tự cài đặt); `joblib` để lưu mô hình.")

# ================================================================== CHƯƠNG 2
chapter(2, "THU THẬP DỮ LIỆU")
h2("2.1. Nguồn dữ liệu và phương pháp thu thập")
para("Dữ liệu được thu thập từ chuyên mục mua bán ô tô của Chợ Tốt (`https://xe.chotot.com/mua-ban-oto`) bằng notebook `crawl_data.ipynb`, hoàn tất ngày **25/10/2026**. Quy trình gồm hai tầng:")
numbered([
    "**Duyệt trang danh sách:** lần lượt truy cập 750 trang danh sách tin mới nhất (`?page=1` đến `?page=750`), lấy đường dẫn của từng tin.",
    "**Đọc trang chi tiết:** với mỗi tin, tải trang và đọc các thẻ HTML có thuộc tính `itemprop` tương ứng với từng thuộc tính của xe (riêng giá đọc từ thẻ `name=\"price\"`). Kết quả được ghi nối tiếp vào `raw_data_oto.csv`.",
])
para("Code có cơ chế **chạy tiếp**: các URL đã có trong file được bỏ qua, nên nếu bị ngắt giữa chừng chỉ cần chạy lại. Các tin \"ưu tiên\" lặp lại ở đầu mọi trang cũng chỉ được ghi một lần. Kết quả là **14.835 tin** với 14 cột (số thứ tự, URL và 12 thuộc tính).")

h2("2.2. Thu thập có trách nhiệm")
bullet("Chỉ lấy thông tin công khai trên trang tin rao, không đăng nhập, không lấy thông tin liên hệ của người bán; dữ liệu chỉ dùng cho mục đích học tập.")
bullet("Nghỉ giữa các yêu cầu (0,5 giây giữa các tin, 1 giây giữa các trang) và tự lùi thời gian khi máy chủ trả về mã 403/429, để không gây tải cho máy chủ.")
bullet("Hạn chế kỹ thuật: danh sách tin được tìm bằng một class CSS sinh tự động; khi trang đổi giao diện cần sửa lại bộ chọn.")

h2("2.3. Từ điển dữ liệu")
para("Bảng 2.1 mô tả các cột thu được. Cột *\"Người bán nhập thế nào\"* rất quan trọng cho bước tiền xử lý: cột **chọn từ danh sách** có tập giá trị cố định, viết thống nhất; còn cột **gõ tự do** (số km, giá) là nguồn lỗi nhập liệu chính.")
table(["Cột", "Ý nghĩa", "Người bán nhập", "Sau xử lý"], [
    ["`mileage_v2`", "Số km đã đi", "**gõ tự do** (dạng 109.366)", "số (`mileage_km`)"],
    ["`carorigin`", "Xuất xứ", "chọn từ danh sách, không bắt buộc", "phân loại"],
    ["`condition_ad`", "Tình trạng (Mới / Đã sử dụng)", "chọn từ danh sách", "phân loại"],
    ["`veh_warranty_policy`", "Chính sách bảo hành", "(một giá trị duy nhất)", "bỏ"],
    ["`carbrand`", "Hãng xe", "chọn từ danh sách", "phân loại"],
    ["`carmodel`", "Dòng xe", "chọn theo hãng", "phân loại"],
    ["`mfdate`", "Năm sản xuất", "chọn (có mục \"trước năm 1980\")", "số"],
    ["`gearbox`", "Hộp số", "chọn (3 giá trị)", "phân loại"],
    ["`fuel`", "Nhiên liệu", "chọn (4 giá trị)", "phân loại"],
    ["`cartype`", "Kiểu dáng", "chọn (9 giá trị), không bắt buộc", "phân loại"],
    ["`carseats`", "Số chỗ ngồi", "chọn (có mục \"Khác\"), không bắt buộc", "số"],
    ["`price`", "Giá rao bán (đồng)", "**gõ tự do**", "biến mục tiêu"],
    ["`url`", "Đường dẫn tin", "sinh tự động, chứa quận/huyện và tỉnh", "trích `province`, rồi bỏ"],
], [3.6, 4.2, 5.0, 3.2], caption="Từ điển dữ liệu thô", align=["l", "l", "l", "l"])

h2("2.4. Cơ chế thu thập và hệ quả cho tiền xử lý")
para("Hiểu **cách dữ liệu được tạo ra** giúp biết mỗi loại lỗi đến từ đâu và nên xử lý thế nào:")
table(["Cơ chế", "Hệ quả", "Cách xử lý"], [
    ["Không tìm thấy thẻ của một thuộc tính → ghi chuỗi `\"N/A\"`", "`N/A` là thiếu thật; pandas mặc định tự đổi thành NaN, làm mất dấu vết", "Đọc với `keep_default_na=False`, rồi tự đổi về NaN"],
    ["Trang hiển thị `\"Đang cập nhật\"` khi người bán không chọn xuất xứ", "Cũng là thiếu thật", "Đổi về NaN"],
    ["Danh sách có sẵn mục `\"Khác\"`, `\"Dòng khác\"`, `\"Nước khác\"`", "Là lựa chọn thật của người bán", "Giữ làm nhãn (trừ `carseats = \"Khác\"`)"],
    ["Số km và giá là ô gõ tự do", "Gõ thiếu/thừa số 0, ghi tiền trả trước, giá giả 123.456.789", "Quy tắc kiểm tra ở mục 3.4, 3.5"],
    ["Xe \"Mới\" không có ô số km", "Số km thiếu ở mọi xe mới", "Điền 0 km"],
    ["Bỏ qua URL đã gặp", "Mỗi dòng là một tin khác nhau; dòng trùng nội dung là tin đăng lại", "Loại trùng theo nội dung"],
    ["URL dạng `mua-ban-oto-<quận>-<tỉnh>/<mã>.htm`", "Có thông tin vị trí", "Trích tỉnh/thành"],
    ["Một lần thu thập, 750 trang mới nhất", "Ảnh chụp thị trường tại một thời điểm; giá là giá rao", "Tuổi xe tính theo năm 2026"],
], [5.0, 5.6, 5.4], caption="Cơ chế thu thập dữ liệu và cách xử lý tương ứng", align=["l", "l", "l"], font=11)

# ================================================================== CHƯƠNG 3
chapter(3, "KHÁM PHÁ VÀ TIỀN XỬ LÝ DỮ LIỆU")
para("Chương này trình bày các bước biến dữ liệu thô thành tập dữ liệu sạch cho mô hình (notebook `preprocess_data.ipynb`). Mỗi quyết định đều kèm lý do, và đều xuất phát từ việc **giá là biến cần dự đoán**.")

h2("3.1. Khám phá dữ liệu ban đầu")
para("Dữ liệu thô có 14.835 dòng và 14 cột. Thị trường tập trung vào các hãng phổ thông: Toyota chiếm khoảng 20% số tin, tiếp theo là Ford, VinFast, Kia, Hyundai (Hình 3.1).")
figure("v05_duoi_dai_hang_dong.png", 15, "Số tin theo hãng và dòng xe, xếp giảm dần: đuôi dài nhãn hiếm (dữ liệu thô)")
para("Giá rao **lệch phải rất mạnh** (độ lệch skewness = 5,08): đa số xe có giá 200-800 triệu, nhưng có một đuôi dài các xe vài tỷ đến hàng chục tỷ đồng . Đặc điểm này ảnh hưởng trực tiếp đến cách chọn biến mục tiêu (mục 3.7).")
para("Hai yếu tố có quan hệ rõ nhất với giá là **năm sản xuất** và **số km**: xe càng mới giá càng cao, xe chạy càng nhiều km giá càng thấp, và quan hệ không tuyến tính (xe mất giá nhanh trong những năm đầu). Giá cũng chênh lệch lớn giữa các hãng.")
figure("p05_eda_gia_theo_tuoi.png", 13, "Giá rao theo tuổi xe (đường cam: giá trung vị)")
figure("p06_eda_gia_theo_hang.png", 13, "Giá rao trung vị của 15 hãng nhiều tin nhất")
figure("p07_eda_gia_theo_km.png", 13, "Giá rao theo số km đã đi (xe đã dùng)")

h2("3.2. Chuẩn hóa kiểu dữ liệu và giá trị giả")
para("Bảng 3.1 thống kê các giá trị \"giả\" trong dữ liệu thô, tức các chuỗi không phải giá trị thật của thuộc tính:")
table(["Cột", "`N/A` (không có thẻ)", "`Đang cập nhật`", "`Khác`", "Số giá trị khác nhau"], [
    ["mileage_v2", "1.811", "0", "0", "1.995"],
    ["carorigin", "0", "6.012", "0", "11"],
    ["veh_warranty_policy", "0", "0", "0", "1"],
    ["cartype", "4.883", "0", "0", "10"],
    ["carseats", "3.028", "0", "659", "12"],
    ["price", "1", "0", "0", "1.433"],
], [4.0, 3.2, 3.0, 2.0, 3.8], caption="Các giá trị giả trong dữ liệu thô")
bullet("`N/A` và `Đang cập nhật` đều là người đăng **không điền thông tin**, nên được đổi thành giá trị thiếu (NaN). Nếu để nguyên, mô hình sẽ coi \"Đang cập nhật\" là một xuất xứ thật.")
bullet("`Khác` trong `carseats` không đổi ra số được nên coi là thiếu; còn `Dòng khác`, `Kiểu dáng khác` là nhãn có nghĩa nên giữ nguyên.")
bullet("`price`, `mileage_v2`, `carseats` dạng chuỗi `1.095.000.000` được bỏ dấu chấm rồi ép sang số; giá đổi sang đơn vị triệu đồng (`price_tr`).")
bullet("`mfdate = \"trước năm 1980\"` (12 dòng) được gán 1979.")

h2("3.3. Trích vị trí, loại cột thừa và tin trùng")
para("**Vị trí từ URL.** Mọi URL có dạng `mua-ban-oto-<quận/huyện>-<tỉnh>/<mã tin>.htm`. So khớp phần cuối với danh sách 63 tỉnh/thành để lấy cột `province` (giữ dạng slug của Chợ Tốt như `tp-ho-chi-minh`, `ha-noi`). Trích được tỉnh cho 100% số tin. Không tách quận/huyện vì có 453 quận/huyện, trung vị chỉ 4 tin mỗi quận, 315 quận dưới 10 tin: dữ liệu quá thưa. Ngay cả ở mức tỉnh, TP.HCM và Hà Nội đã chiếm 67% số tin, 28 tỉnh có dưới 30 tin.")
para("**Cột loại bỏ:** `no` và `url` là định danh, không mang thông tin dự đoán; `veh_warranty_policy` chỉ có một giá trị duy nhất (\"Bảo hành hãng\") ở mọi dòng.")
para("**Tin trùng:** 691 dòng (4,7%) giống hệt nhau ở mọi thuộc tính và giá (331 trong số đó là xe mới, thường do đại lý đăng nhiều xe cùng cấu hình). Giữ lại sẽ làm lệch phân phối, và khi chia train/test thì cùng một xe có thể nằm ở cả hai phía, làm điểm đánh giá cao ảo. Chỉ giữ một bản, còn lại **14.144 dòng**.")

h2("3.4. Giá trị bất hợp lý")
para("Bước này làm **trước** khi điền thiếu, để không điền dựa trên số liệu sai. Phân biệt hai loại: **sai dữ liệu** (nhập nhầm, không thể xảy ra) thì loại hoặc sửa; **ngoại lệ thật** (xe siêu sang giá chục tỷ) thì giữ, vì đó là xe có thật.")
table(["Cột", "Quy tắc", "Lý do", "Số dòng"], [
    ["Giá", "Loại nếu < 20 triệu hoặc > 10 tỷ", "Gõ thiếu/thừa số 0; không biết giá đúng nên loại", "64"],
    ["Giá", "Loại giá là dãy số liên tiếp (123.456.789...)", "Giá giả khi người bán không muốn công khai giá", "14"],
    ["Tình trạng", "\"Mới\" nhưng hơn 3 năm tuổi → \"Đã sử dụng\"", "Ghi sai tình trạng (ví dụ Tucson 2011 \"mới\")", "6 (sửa)"],
    ["Số km", "Xe đã dùng > 500.000 km → thiếu", "Nhập thừa số 0", "106"],
    ["Số km", "Xe ≥ 2 năm tuổi mà < 100 km → thiếu", "Nhập theo đơn vị nghìn km", "447"],
    ["Số km", "> 100.000 km/năm → thiếu", "Bất khả thi với xe dân dụng", "25"],
], [2.2, 5.4, 6.4, 2.0], caption="Quy tắc phát hiện giá trị bất hợp lý", align=["l", "l", "l", "c"], font=11)
para("Các giá như 399.999.999 hay 888.888.888 vẫn được **giữ**, vì đó là cách định giá tâm lý hoặc chọn số đẹp, vẫn là giá thật. Sau bước này còn **14.066 dòng** (loại 78 dòng); 578 giá trị số km bất hợp lý được chuyển thành thiếu để điền lại ở mục 3.6.")

h2("3.5. Giá trả trước / đặt cọc")
para("Ngưỡng 20 triệu chỉ bắt được lỗi gõ thiếu số. Vẫn còn những tin mà \"giá\" thực ra là **số tiền trả trước hoặc đặt cọc**, ví dụ xe đời 2025-2026 còn mới mà giá chỉ 30-100 triệu. Một ngưỡng tuyệt đối không bắt được trường hợp này, vì 100 triệu hợp lý với xe 20 năm tuổi nhưng vô lý với xe 1 năm tuổi.")
para("**Quy tắc:** loại tin có giá **thấp hơn 30% giá trung vị của các xe cùng hãng, dòng và năm sản xuất** (chỉ xét nhóm có ít nhất 3 xe). Không xe nào mất 70% giá so với xe cùng đời chỉ vì tình trạng hay số km. Kết quả: loại **129 tin**, 94% trong số đó có giá dưới 200 triệu; còn lại **13.937 dòng**.")
table(["Xe", "Năm", "Tình trạng", "Giá rao (triệu)", "Trung vị nhóm (triệu)"], [
    ["Toyota Camry", "2026", "Đã sử dụng", "300", "1.400"],
    ["Ford Everest", "2026", "Đã sử dụng", "320", "1.325"],
    ["Toyota Fortuner", "2026", "Mới", "125", "1.249"],
    ["Hyundai Santa Fe", "2025", "Mới", "220", "959"],
    ["Mercedes Benz E Class", "2019", "Đã sử dụng", "156", "889"],
], [5.0, 1.6, 3.0, 3.0, 3.4], caption="Ví dụ tin bị loại vì nghi ghi giá trả trước / đặt cọc")
para("Đây là bước **sửa lỗi nhãn**: quy tắc dùng giá trên toàn bộ dữ liệu nhưng chỉ để loại dòng có đáp án sai, không đưa thông tin giá vào đặc trưng nên không gây rò rỉ dữ liệu vào đầu vào của mô hình.")
para("**Ngoại lệ theo IQR** của log giá còn 696 dòng (5,0%), chủ yếu là xe sang thật (Mercedes, Lexus, Bentley, Range Rover...) hoặc xe cổ rất rẻ. Đây là thông tin thật của thị trường nên **giữ nguyên**; ảnh hưởng của chúng được giảm bằng biến đổi log.")

h2("3.6. Xử lý giá trị thiếu")
para("Trước khi điền cần hiểu **vì sao thiếu**, vì cách điền phụ thuộc cơ chế thiếu.")
figure("p02_thieu_va_cach_xu_ly.png", 16, "Tỉ lệ giá trị thiếu của từng cột và cách xử lý")
table(["Cột", "Cơ chế thiếu", "Cách xử lý"], [
    ["`mileage_km`", "**Toàn bộ** 1.235 xe mới thiếu km (trang tin xe mới không có ô số km); chỉ 244 xe đã dùng thiếu thật, cộng 578 giá trị bất hợp lý", "Xe mới: **0 km**. Xe đã dùng: trung vị theo năm sản xuất (814 xe, 5,8%), đánh dấu `mileage_missing = 1`"],
    ["`carorigin`", "Thiếu ~41%; một mẫu xe có thể vừa nhập khẩu vừa lắp ráp", "Nhãn riêng **\"Không rõ\"** (5.787 dòng)"],
    ["`cartype`", "Thiếu ~33%; kiểu dáng gần như do dòng xe quyết định", "Mode theo hãng + dòng xe (4.336); còn lại \"Không rõ\" (243)"],
    ["`carseats`", "Thiếu ~25%; số chỗ do dòng xe quyết định", "Mode theo dòng xe (3.352) → theo hãng (183) → mode chung (4)"],
    ["`price`", "1 dòng", "**Xóa** (không bịa biến mục tiêu)"],
], [2.6, 6.8, 6.6], caption="Cơ chế thiếu và cách điền cho từng cột", align=["l", "l", "l"], font=11)
para("**Vì sao không điền trung vị số km cho xe mới?** Điền trung vị sẽ biến xe mới thành \"xe đã chạy 66.000 km\", sai lệch nghiêm trọng. Đây là thiếu **có cấu trúc** (không đo vì không có ý nghĩa), giá trị đúng là 0.")
para("**Cột đánh dấu `mileage_missing`:** giá trị điền cho xe đã dùng chỉ là \"đoán trung bình\". Cột này cho mô hình biết dòng nào là ước lượng để tự học mức độ tin cậy; việc người bán không khai số km cũng có thể là một tín hiệu về giá.")
para("**Vì sao `cartype` không lùi về cấp hãng còn `carseats` thì có?** Trong một dòng xe, kiểu dáng gần như cố định; nhưng một hãng bán đủ loại (Ford có Ranger bán tải, Everest SUV, Focus sedan), nên gán kiểu dáng phổ biến của hãng là đoán mò. Số chỗ thì đa số là 5, đoán sai theo hãng thường chỉ lệch 5 ↔ 7 chỗ.")

h2("3.7. Biến đổi log cho giá")
para("Giá lệch phải mạnh khiến mô hình hồi quy bị kéo bởi vài xe rất đắt. Lấy logarit đưa phân phối về gần đối xứng (độ lệch giảm từ **5,08 xuống −0,20**) và biến chênh lệch tuyệt đối thành chênh lệch **tương đối**: xe 100 → 150 triệu và xe 1 → 1,5 tỷ đều là tăng 50%, khớp với cách người mua nghĩ về giá và với thước đo MAPE.")
equation(r"y = \ln(p)")
equation(r"\hat{p} = e^{\hat{y}}")
para("Với p là giá (triệu đồng), mô hình học `price_log`, dự đoán xong dùng `exp()` để đổi về triệu đồng; **mọi thước đo đều tính trên giá thật**. Lựa chọn này được kiểm chứng bằng thí nghiệm ở mục 5.8.")
figure("p03_gia_va_log_gia.png", 15, "Phân bố giá gốc và log giá sau tiền xử lý")

h2("3.8. Gộp nhãn hiếm theo cấp")
para("Dữ liệu có 69 hãng và 509 tổ hợp hãng + dòng xe, nhiều nhãn chỉ có 1-5 tin. One-hot nguyên vẹn sinh hàng trăm cột thưa, mô hình học quy luật từ vài mẫu là học nhiễu. Cách gộp đơn giản (mọi dòng hiếm thành một nhãn `\"Dòng khác\"`) có vấn đề: nhãn này chiếm 17% dữ liệu nhưng tới **65% số xe trên 2 tỷ**, một chiếc Mercedes G63 và một chiếc Kia đời cũ mang cùng nhãn.")
para("Đề tài dùng cách **gộp theo cấp**: mỗi xe lùi về nhóm gần nhất còn đủ mẫu.")
table(["Cấp", "Quy tắc", "Ví dụ"], [
    ["Dòng xe", "Giữ `Hãng + Dòng` nếu có ≥ 10 tin", "Toyota Vios"],
    ["→ Hãng", "Dòng hiếm → `\"<Hãng> - dòng khác\"`", "Mercedes Benz - dòng khác"],
    ["Hãng", "Giữ hãng nếu có ≥ 30 tin", "Volvo, Subaru"],
    ["→ Nhóm hãng", "Hãng hiếm → `\"Hãng cao cấp khác\"` / `\"Hãng phổ thông khác\"`", "Bentley → cao cấp; Wuling → phổ thông"],
], [3.0, 7.6, 5.4], caption="Quy tắc gộp nhãn hiếm theo cấp", align=["l", "l", "l"])
para("Nhóm cao cấp / phổ thông được phân theo **định vị thương hiệu** (hiểu biết thị trường), **không** dựa vào giá trong dữ liệu, để không dùng biến mục tiêu khi tạo đặc trưng. Kết quả: 69 hãng → **28 nhóm**, 509 dòng → **214 nhóm**; 8,9% số xe mang nhãn \"... - dòng khác\". Nhóm \"Hãng cao cấp khác\" có 93 tin từ 13 hãng (giá trung vị 1 tỷ), \"Hãng phổ thông khác\" có 188 tin từ 30 hãng (giá trung vị 390 triệu).")
para("Cách gộp và ngưỡng được **chọn bằng 5-fold CV chỉ trên tập train** (Bảng 3.6). Gộp theo cấp cải thiện rõ Ridge (MAPE 21,9% → 18,9%; xe > 2 tỷ 43,0% → 32,2%) vì mô hình tuyến tính chỉ có một hệ số cho mỗi nhãn. Random Forest được lợi ít hơn vì tự bù thông tin bằng các đặc trưng khác. Không phương án nào làm mô hình xấu đi.")
table(["Phương án", "Ridge: MAPE", "< 200 tr", "> 2 tỷ", "RF: MAPE", "< 200 tr", "> 2 tỷ"], [
    ["Một nhãn chung (50 / 30)", "0,219", "0,522", "0,430", "0,163", "0,374", "0,292"],
    ["Gộp theo cấp (50 / 30)", "0,206", "0,474", "0,392", "0,163", "0,361", "0,278"],
    ["Gộp theo cấp (30 / 10) – chọn", "0,189", "0,449", "0,322", "0,161", "0,359", "0,287"],
], [5.0, 1.9, 1.8, 1.8, 1.9, 1.8, 1.8], caption="So sánh các cách gộp nhãn (5-fold CV trên train, MAPE)", bold_rows=(2,), font=11)

h2("3.9. Tạo đặc trưng")
bullet("`age` = 2026 − năm sản xuất: tuổi xe có ý nghĩa trực tiếp hơn năm tuyệt đối.")
bullet("`km_per_year` = km / max(tuổi, 1): xe 5 năm đi 150.000 km (nhiều khả năng chạy dịch vụ) khác hẳn xe 5 năm đi 50.000 km.")
bullet("`mileage_missing`: 1 nếu số km là giá trị điền.")
bullet("`province`: tỉnh/thành của tin đăng.")
para("**Vị trí có giúp dự đoán không?** Thử thêm miền (3 nhóm) hoặc tỉnh/thành (15 nhóm, tỉnh dưới 100 tin gộp theo miền) bằng cùng cách CV: MAPE chỉ chênh ở chữ số thứ ba (Ridge 0,1895 / 0,1894 / 0,1897; Random Forest 0,1610 / 0,1604 / 0,1611), nhỏ hơn mức dao động giữa các lần chia fold. Giá rao chủ yếu do bản thân chiếc xe quyết định; chênh lệch giá giữa các miền chủ yếu do **cơ cấu xe khác nhau**, mà mô hình đã thấy qua các đặc trưng khác. Vì vậy `province` được giữ trong file để mô tả dữ liệu nhưng **không đưa vào mô hình**.")

h2("3.10. Dữ liệu mất cân bằng")
para("Biến mục tiêu là số liên tục nên không có \"lớp thiểu số\" như bài toán phân loại. Mất cân bằng ở đây thể hiện ở **mật độ dữ liệu không đều theo vùng giá**: phần lớn xe ở phân khúc phổ thông, xe sang vài tỷ rất ít. Mô hình tối ưu sai số trung bình sẽ học tốt vùng đông mẫu và kém ở vùng hiếm. Để định lượng, giá được chia thành 5 phân khúc (chỉ dùng để phân tích, chia mẫu và đánh giá, không làm đặc trưng):")
table(["Phân khúc", "Số tin", "Tỉ lệ", "Ít hơn phân khúc lớn nhất"], [
    ["< 200 triệu", "1.972", "14,1%", "3,0 lần"],
    ["200-500 triệu", "5.907", "42,4%", "1,0 lần"],
    ["500-1.000 triệu", "4.529", "32,5%", "1,3 lần"],
    ["1-2 tỷ", "1.071", "7,7%", "5,5 lần"],
    ["> 2 tỷ", "458", "3,3%", "12,9 lần"],
], [4.0, 3.0, 3.0, 5.0], caption="Phân bố dữ liệu theo phân khúc giá")
figure("p04_phan_khuc_gia.png", 13, "Số tin đăng theo phân khúc giá")
para("Phân khúc > 2 tỷ ít hơn phân khúc lớn nhất gần 13 lần: mất cân bằng thật sự. Các đặc trưng cũng lệch: xe mới chỉ 8,5%, Hybrid 2,2%, hộp số bán tự động 0,7%. Các biện pháp áp dụng:")
bullet("**Chia train/test phân tầng theo phân khúc giá**, để tập test có đủ xe ở mọi phân khúc.")
bullet("**Học trên log giá**, giảm ảnh hưởng của đuôi phải.")
bullet("**Báo cáo MAPE theo từng phân khúc**, không chỉ một con số trung bình.")
bullet("**Thử trọng số mẫu, Random Oversampling, Random Undersampling** bằng CV trên tập train (mục 4.7 và 5.8).")
para("File `clean_data_oto.csv` **chỉ chứa dữ liệu thật**: nếu ghi mẫu nhân bản vào file, người dùng sau chia ngẫu nhiên sẽ khiến bản sao của cùng một xe nằm ở cả train lẫn test.")

h2("3.11. Tổng kết tiền xử lý")
table(["Bước", "Số dòng còn lại", "Thay đổi"], [
    ["Dữ liệu thô", "14.835", ""],
    ["Loại tin trùng nội dung", "14.144", "−691"],
    ["Loại giá thiếu / < 20 triệu / > 10 tỷ / giá giả", "14.066", "−78"],
    ["Loại giá trả trước / đặt cọc", "13.937", "−129"],
    ["**Dữ liệu sạch** (giữ 93,9%)", "**13.937**", "18 cột, không còn giá trị thiếu"],
], [8.0, 3.5, 4.5], caption="Số dòng qua từng bước tiền xử lý", bold_rows=(4,))
figure("p01_pheu_lam_sach.png", 15, "Số tin còn lại qua từng bước làm sạch")
para("File `clean_data_oto.csv` gồm 18 cột: `carbrand`, `carbrand_grp`, `carmodel`, `carmodel_grp`, `carorigin`, `condition_ad`, `gearbox`, `fuel`, `cartype`, `carseats`, `mfdate`, `age`, `mileage_km`, `mileage_missing`, `km_per_year`, `province`, `price_tr`, `price_log`. Cột gốc `carbrand`, `carmodel` được giữ để tra cứu; mô hình dùng cột đã gộp `*_grp`.")

# ================================================================== CHƯƠNG 4
chapter(4, "XÂY DỰNG MÔ HÌNH RANDOM FOREST")
para("Chương này trình bày mô hình chính của đề tài, **Random Forest**, với hai bản: bản dùng thư viện `scikit-learn` và bản **tự cài đặt bằng NumPy**. **Ridge Regression** chỉ đóng vai trò mô hình so sánh (mục 4.6), dùng thẳng bản thư viện. Toàn bộ nằm trong notebook `model.ipynb`.")

h2("4.1. Thiết kế thí nghiệm")
h3("4.1.1. Chia dữ liệu")
para("Dữ liệu được chia **80% train (11.149 tin) / 20% test (2.788 tin)**, phân tầng theo phân khúc giá, `random_state = 42`. Phân tầng đảm bảo tập test có đủ xe ở mọi phân khúc: nếu chia ngẫu nhiên thông thường, tập test có thể gần như không có xe trên 2 tỷ (chỉ khoảng 3% dữ liệu), và điểm đánh giá sẽ đẹp nhưng không phản ánh khả năng với xe sang (Hình 4.1).")
figure("m01_chia_train_test.png", 13, "Chia train/test phân tầng: tỉ lệ các phân khúc giá như nhau ở hai tập")
para("**Không tách tập validation riêng.** Tham số được chọn bằng **cross-validation** (CV) hoặc **out-of-bag** (OOB) **trên tập train**. Với chỉ khoảng 460 xe trên 2 tỷ, một tập validation cố định quá nhỏ và dễ dao động; CV dùng lần lượt mọi phần của train làm validation nên ổn định hơn. Tập test chỉ dùng **một lần** để báo cáo kết quả cuối.")
para("**Tính tái lập.** Mọi bước ngẫu nhiên (chia dữ liệu, chia fold, bootstrap của rừng, chọn đặc trưng ngẫu nhiên, bootstrap khoảng tin cậy) đều cố định `random_state = 42`. Chạy lại notebook cho đúng cùng mô hình và cùng số liệu; chỉ thời gian huấn luyện thay đổi theo tải của máy.")

h3("4.1.2. Biến mục tiêu và đặc trưng")
para("Mô hình học **ln(giá)** (mục 3.7), dự đoán xong dùng `exp()` để đổi về triệu đồng, mọi thước đo tính trên giá thật. Đặc trưng chỉ gồm thông tin người mua thấy trong tin đăng:")
table(["Nhóm", "Random Forest (mô hình chính)", "Ridge (so sánh)"], [
    ["Phân loại", "hãng, dòng xe (đã gộp nhãn), xuất xứ, tình trạng, hộp số, nhiên liệu, kiểu dáng; **one-hot** (mục 4.2)", "như RF, one-hot"],
    ["Tuổi xe", "`age`", "`age`, `age²`"],
    ["Số km", "`mileage_km`, `km_per_year` giữ nguyên", "`log(1+km)`, `log(1+km/năm)`"],
    ["Khác", "`carseats`, `mileage_missing`", "như RF"],
    ["Chuẩn hóa", "không cần", "StandardScaler"],
    ["Số cột sau mã hóa", "277", "278"],
], [3.0, 7.6, 5.4], caption="Đặc trưng đầu vào của hai mô hình", align=["l", "l", "l"], font=11)
bullet("**Random Forest không cần biến đổi đặc trưng số:** cây chỉ so sánh ngưỡng (\"tuổi ≤ 5?\"), nên mọi biến đổi đơn điệu như `age²` hay `log` cho đúng cùng một phép chia. Thêm `age²` không làm mô hình tốt hơn, chỉ làm độ quan trọng của tuổi xe bị chia đôi cho hai cột. Cây cũng không phụ thuộc thang đo nên không chuẩn hóa.")
bullet("**Ridge cần `age²`, `log` và chuẩn hóa** vì chỉ cộng các hiệu ứng tuyến tính, trong khi xe mất giá nhanh những năm đầu rồi chậm lại.")
bullet("**Không dùng:** `price_tr`, `price_log` (đáp án), `mfdate` (trùng `age`), `province` (không giúp dự đoán, mục 3.9). Bộ mã hóa **fit trên train** rồi áp lên test.")

h3("4.1.3. Thước đo đánh giá")
para("Tất cả thước đo tính trên **giá thật (triệu đồng)**:")
equation(r"\mathrm{MAPE} = \frac{1}{n}\sum_{i=1}^{n}\frac{|p_i - \hat{p}_i|}{p_i}")
equation(r"\mathrm{MAE} = \frac{1}{n}\sum_{i=1}^{n}|p_i - \hat{p}_i| \qquad \mathrm{RMSE} = \sqrt{\frac{1}{n}\sum_{i=1}^{n}(p_i - \hat{p}_i)^2}")
para("trong đó pᵢ là giá thật và p̂ᵢ là giá dự đoán của xe thứ i (triệu đồng). Bảng 4.2 tóm tắt ý nghĩa, chiều tốt và mức tham khảo của từng thước đo:", indent=False)
table(["Thước đo", "Ý nghĩa", "Tốt khi", "Mức tham khảo"], [
    ["**MAPE**", "Sai số phần trăm trung bình. **Thước đo chính**: lệch 100 triệu là nhỏ với xe 3 tỷ nhưng rất lớn với xe 150 triệu", "**thấp**", "< 10%: rất tốt; 10-20%: tốt; 20-50%: tạm được; > 50%: kém"],
    ["MedAPE", "Trung vị sai số phần trăm: sai số của một xe \"điển hình\", không bị vài tin sai nặng kéo lệch", "**thấp**", "cùng thang với MAPE"],
    ["MAE", "Trung bình đoán lệch bao nhiêu triệu đồng", "**thấp**", "không có ngưỡng chung, tùy mức giá"],
    ["RMSE", "Như MAE nhưng phạt nặng lỗi lớn; RMSE lớn hơn MAE nhiều nghĩa là có vài lỗi rất lớn", "**thấp**", "luôn ≥ MAE"],
    ["R² (log)", "Tỉ lệ biến thiên của log giá được mô hình giải thích", "**cao**", "1: hoàn hảo; 0: không hơn đoán trung bình; < 0: kém hơn"],
    ["Lệch ≤ 10% / ≤ 20%", "Tỉ lệ xe có giá dự đoán nằm trong ±10% / ±20% giá thật", "**cao**", "so với baseline"],
], [2.6, 7.0, 1.6, 4.8], caption="Các thước đo đánh giá và cách đọc", align=["l", "l", "c", "l"], font=10.5)
para("Thang MAPE ở cột cuối là quy ước thường dùng trong dự báo [9], không phải tiêu chuẩn cứng. Không chỉ số nào \"tốt\" khi đứng một mình: luôn so với **mô hình mốc (baseline)** luôn đoán giá trung vị của tập train. Ngoài ra luôn báo **MAPE theo từng phân khúc giá**, vì một con số trung bình che mất việc mô hình kém ở xe rất rẻ hoặc rất đắt.")

h2("4.2. Chuẩn bị đặc trưng: one-hot hay mã số thứ tự?")
para("Biến phân loại phải đổi thành số trước khi đưa vào cây. Có hai cách:")
table(["Cách", "Làm gì", "Ưu điểm", "Nhược điểm"], [
    ["**One-hot**", "Mỗi nhãn thành một cột 0/1 (dòng xe có 214 nhãn nên sinh 214 cột)", "Không tạo thứ tự giả; một phép chia tách thẳng một nhãn khỏi phần còn lại", "Nhiều cột (277); độ quan trọng của một đặc trưng bị chia nhỏ"],
    ["**Mã số thứ tự**", "Mỗi nhãn thành một số nguyên 0, 1, 2... trong **một** cột", "Chỉ 12 cột", "Thứ tự nhãn tùy ý (theo bảng chữ cái); phải chia nhiều lần mới tách được một nhãn"],
], [2.6, 4.6, 4.6, 4.2], caption="Hai cách mã hóa biến phân loại cho cây", align=["l", "l", "l", "l"], font=11)
para("Với mô hình tuyến tính chỉ có one-hot là đúng, nhưng cây không bị thứ tự giả ảnh hưởng nặng, nên đây là **câu hỏi thực nghiệm**. So sánh bằng **5-fold CV phân tầng chỉ trên tập train** (bộ mã hóa fit trong phần train của từng fold), rừng 150 cây, `max_depth = 30`; riêng `max_features` thử vài mức với mã số thứ tự vì số cột khác hẳn:")
table(["Phương án", "Số cột", "MAPE", "< 200 tr", "200-500", "500-1.000", "1-2 tỷ", "> 2 tỷ"], [
    ["**One-hot**, max_features = 0,33", "277", "**16,4%**", "36,5%", "12,8%", "10,8%", "16,7%", "28,9%"],
    ["Mã số, max_features = 0,33", "12", "18,4%", "40,5%", "14,3%", "11,7%", "18,8%", "39,0%"],
    ["Mã số, max_features = 0,6", "12", "16,7%", "37,8%", "13,1%", "10,5%", "17,8%", "30,1%"],
    ["Mã số, max_features = 1,0", "12", "16,9%", "38,3%", "13,3%", "10,8%", "17,7%", "29,0%"],
], [4.6, 1.3, 1.5, 1.6, 1.6, 1.7, 1.6, 1.6], caption="So sánh cách mã hóa (5-fold CV trên train, MAPE)", bold_rows=(0,), font=10.5)
figure("m16_ma_hoa.png", 16, "MAPE tổng và MAPE xe trên 2 tỷ theo cách mã hóa")
para("**Kết luận: giữ one-hot.** One-hot cho MAPE thấp nhất, và chênh rõ nhất ở hai đầu thang giá. Lý do: với one-hot, **một** phép chia tách thẳng một dòng xe (ví dụ Mercedes G Class) khỏi phần còn lại; với mã số thứ tự, các nhãn xếp theo bảng chữ cái nên dòng xe sang và xe rẻ nằm xen kẽ, cây phải chia nhiều lần mới cô lập được một dòng xe, và dòng xe hiếm thì không đủ mẫu để chia tới đó. One-hot cũng hợp với bản tự cài đặt: mỗi cột 0/1 nằm gọn trong 2 bin, còn mã số của hơn 200 dòng xe sẽ bị dồn vào 32 bin.")

h2("4.3. Cơ sở lý thuyết của Random Forest")
h3("4.3.1. Thành phần: cây hồi quy CART")
para("Cây quyết định chia dữ liệu bằng các câu hỏi có/không, ví dụ \"tuổi xe ≤ 13,5?\", \"hộp số là số sàn?\". Tại mỗi nút, thuật toán thử các đặc trưng và ngưỡng, chọn phép chia làm **giảm tổng bình phương sai số (SSE)** nhiều nhất. Với nút có n mẫu, tổng nhãn S, chia thành nhánh trái (n_L mẫu, tổng S_L) và nhánh phải (n_R, S_R):")
equation(r"\text{gain} = \frac{S_L^2}{n_L} + \frac{S_R^2}{n_R} - \frac{S^2}{n}")
para("Công thức này là mức giảm SSE khi thay giá trị dự đoán chung của nút bằng trung bình riêng của từng nhánh. Lặp lại cho đến khi đạt độ sâu tối đa hoặc nút quá ít mẫu; giá trị dự đoán ở mỗi lá là **trung bình log giá** của các xe trong lá.")
para("Hình 4.3 là một cây **thật** độ sâu 3 học từ tập train. Câu hỏi đầu tiên cây chọn là tuổi xe (≤ 13,5 năm), tiếp theo là hộp số số sàn, rồi kiểu dáng và số chỗ. Chỉ với 3 tầng câu hỏi, cây đã tách được nhóm xe khoảng 617 triệu (xe đời mới, số tự động, không phải hatchback) khỏi nhóm khoảng 73 triệu (xe cũ hơn 20 năm, số sàn). Cây thật trong rừng sâu tới 30 tầng và có hàng nghìn nút.")
figure("m17_cay_minh_hoa.png", 16, "Một cây hồi quy độ sâu 3 học từ tập train (ô cam: lá, kèm giá dự đoán)")
para("**Vấn đề của một cây:** cây sâu học thuộc dữ liệu train (độ chệch thấp) nhưng **phương sai cao**: đổi vài mẫu là cấu trúc cây thay đổi hẳn, và dự đoán trên dữ liệu mới kém ổn định.")
h3("4.3.2. Rừng ngẫu nhiên")
para("Random Forest [3] giảm phương sai bằng cách lấy trung bình nhiều cây **khác nhau**:")
numbered([
    "**Bootstrap:** mỗi cây học trên một mẫu lấy ngẫu nhiên **có hoàn lại** từ train, cùng kích thước. Do lấy có hoàn lại, mỗi mẫu bootstrap chứa khoảng 63% số xe khác nhau (xác suất một xe không được chọn sau n lần bốc là (1 − 1/n)ⁿ ≈ e⁻¹ ≈ 37%).",
    "**Chọn ngẫu nhiên đặc trưng:** mỗi lần chia nút chỉ xét một tỉ lệ `max_features` đặc trưng ngẫu nhiên, để các cây ít giống nhau.",
    "**Lấy trung bình** dự đoán của tất cả T cây:",
])
equation(r"\hat y(x) = \frac{1}{T}\sum_{t=1}^{T} f_t(x)")
para("Trung bình của nhiều cây **ít tương quan** có phương sai nhỏ hơn hẳn một cây, trong khi độ chệch gần như không đổi. Đó là lý do bước 2 quan trọng: nếu mọi cây đều chọn tuổi xe ở gốc, các cây sẽ giống nhau và lấy trung bình không giúp được nhiều.")
para("**Out-of-bag (OOB):** khoảng 37% số xe không nằm trong mẫu bootstrap của một cây. Dự đoán mỗi xe bằng những cây \"chưa thấy\" nó cho ra một **ước lượng sai số khách quan mà không cần tập validation riêng**.")
table(["Siêu tham số", "Tác dụng"], [
    ["`n_estimators`", "Số cây; nhiều hơn thì ổn định hơn, đến một mức thì bão hòa"],
    ["`max_depth`", "Độ sâu tối đa của cây"],
    ["`min_samples_leaf`", "Số mẫu tối thiểu ở mỗi lá; lớn hơn thì cây \"mượt\" hơn, ít học thuộc hơn"],
    ["`max_features`", "Tỉ lệ đặc trưng xét ở mỗi lần chia; nhỏ hơn thì các cây đa dạng hơn"],
], [4.0, 12.0], caption="Siêu tham số của Random Forest", align=["l", "l"])

h2("4.4. Random Forest: bản thư viện")
para("Dùng `sklearn.ensemble.RandomForestRegressor`, tinh chỉnh bằng **RandomizedSearchCV**: thử ngẫu nhiên 12 cấu hình trong không gian `max_depth` ∈ {20, 30, không giới hạn}, `min_samples_leaf` ∈ {1, 2, 3, 5}, `max_features` ∈ {0,33; 0,5; 0,8}, với 300 cây, đánh giá bằng **3-fold CV phân tầng theo phân khúc giá** trên train (RMSE trên log giá, tương đương sai số tương đối nên nhất quán với MAPE). Tìm kiếm mất khoảng 3,5 phút.")
table(["n_estimators", "min_samples_leaf", "max_features", "max_depth", "CV RMSE (log)"], [
    ["300", "1", "0,50", "30", "0,2692"],
    ["300", "1", "0,33", "30", "0,2701"],
    ["300", "2", "0,80", "30", "0,2763"],
    ["300", "2", "0,50", "30", "0,2766"],
    ["300", "2", "0,33", "30", "0,2793"],
], [3.0, 3.4, 3.0, 2.8, 3.4], caption="Năm cấu hình Random Forest tốt nhất (thư viện, 3-fold CV)", bold_rows=(0,), align=["c"] * 5)
para("Cấu hình tốt nhất: **300 cây, `max_depth = 30`, `min_samples_leaf = 1`, `max_features = 0,5`**; huấn luyện lại trên toàn bộ train mất khoảng 9 giây, **OOB R² (log) = 0,896**. Các cấu hình tốt đều có cây sâu và lá nhỏ: dữ liệu có nhiều tổ hợp hãng-dòng-đời xe riêng biệt nên mô hình cần đủ độ phức tạp để học từng nhóm; việc lấy trung bình 300 cây đã kiểm soát phương sai.")

h2("4.5. Random Forest: bản tự cài đặt bằng NumPy")
h3("4.5.1. Cấu trúc")
table(["Lớp", "Nhiệm vụ"], [
    ["`Binner`", "Rời rạc hóa mỗi đặc trưng thành tối đa 32 bin theo phân vị (fit trên train)"],
    ["`RegressionTreeScratch`", "Cây hồi quy CART: tìm phép chia tốt nhất theo công thức gain, dựng cây bằng ngăn xếp (không đệ quy)"],
    ["`RandomForestScratch`", "Bootstrap + chọn ngẫu nhiên đặc trưng + lấy trung bình; tính OOB và độ quan trọng đặc trưng"],
], [5.0, 11.0], caption="Các lớp của Random Forest tự cài đặt", align=["l", "l"])
para("**Tăng tốc bằng histogram.** Cách cơ bản phải sắp xếp từng cột ở mỗi nút để thử mọi ngưỡng, rất chậm khi viết bằng Python. Thay vào đó, mỗi đặc trưng được chia sẵn thành tối đa 32 bin **một lần** (kỹ thuật của LightGBM [5] và HistGradientBoosting). Tại mỗi nút chỉ cần đếm số mẫu và tổng nhãn trong từng bin (`np.bincount`), rồi cộng dồn (`np.cumsum`) là có n_L, S_L của **mọi ngưỡng cùng lúc**, thay vào công thức gain. Cột one-hot chỉ có 2 giá trị nên không mất thông tin; cột số mất rất ít độ chi tiết.")
para("**Độ quan trọng đặc trưng** = tổng gain mà đặc trưng mang lại ở mọi nút của mọi cây, chuẩn hóa về tổng bằng 1 (cùng định nghĩa với `sklearn`).")
h3("4.5.2. Kiểm chứng thành phần: một cây")
para("Trước khi dựng rừng, so sánh **một cây** tự viết (xét mọi đặc trưng ở mỗi nút) với `DecisionTreeRegressor` khi dùng **cùng tham số** (lá ≥ 5 mẫu, không giới hạn độ sâu):")
table(["Mô hình", "Số nút", "R² (log) train", "R² (log) test", "MAPE test"], [
    ["Một cây – thư viện", "3.481", "0,929", "0,842", "19,6%"],
    ["Một cây – tự cài đặt", "3.429", "0,927", "0,851", "18,6%"],
], [5.0, 2.2, 3.0, 3.0, 2.8], caption="Kiểm chứng cây hồi quy tự cài đặt")
para("Hai cây có số nút và R² gần như nhau, nên **lớp cây tự viết đúng**. Khoảng cách giữa R² train (0,93) và R² test (0,84-0,85) chính là **phương sai cao** của một cây. Bảng này dùng tập test chỉ để kiểm tra tính đúng của code, không để chọn tham số.")
h3("4.5.3. Tinh chỉnh bằng OOB")
para("Bản tự viết chậm hơn thư viện nhiều, nên thay vì CV (huấn luyện 3-5 lần mỗi cấu hình), dùng **OOB RMSE**: mỗi cấu hình chỉ huấn luyện **một lần** với 30 cây. Thử 6 cấu hình trong cùng không gian tham số, mất khoảng 3 phút. OOB dùng được ở đây vì mỗi xe chỉ xuất hiện một lần trong train (mục 4.7 giải thích vì sao OOB **không** dùng được khi oversampling).")
table(["max_depth", "min_samples_leaf", "max_features", "OOB RMSE (log)", "OOB R²"], [
    ["30", "2", "0,50", "0,2704", "0,886"],
    ["30", "2", "0,33", "0,2741", "0,883"],
    ["30", "3", "0,33", "0,2779", "0,879"],
    ["20", "1", "0,33", "0,2938", "0,865"],
    ["20", "2", "0,33", "0,2946", "0,864"],
    ["20", "3", "0,50", "0,2972", "0,862"],
], [2.6, 3.4, 3.0, 3.4, 2.6], caption="Tinh chỉnh Random Forest tự cài đặt bằng OOB (30 cây mỗi cấu hình)", bold_rows=(0,), align=["c"] * 5)
h3("4.5.4. Từ một cây đến một rừng")
para("Với cấu hình tốt nhất, huấn luyện **100 cây** mất khoảng 2,4 phút, **OOB R² (log) = 0,892** (bản thư viện 0,896). Hình 4.4 theo dõi OOB RMSE khi số cây tăng dần: sai số giảm rất nhanh trong vài chục cây đầu (từ 0,346 với 1 cây) rồi **bão hòa** quanh 0,263; thêm cây sau điểm đó chỉ tốn thời gian.")
figure("m05_rf_so_cay.png", 12.5, "OOB RMSE của rừng tự cài đặt theo số cây")
para("Hình 4.5 so sánh một cây và cả rừng trên tập test: lấy trung bình nhiều cây giảm MAPE từ 18,6-19,6% xuống 15,5-15,7% và tăng R² từ khoảng 0,85 lên 0,89. Bản tự cài đặt đi theo đúng xu hướng của bản thư viện ở cả hai mức.")
figure("m18_cay_vs_rung.png", 16, "Một cây so với cả rừng trên tập test (hai bản cài đặt)")
table(["Mô hình", "MAPE", "MedAPE", "R² (log)", "Lệch ≤ 20%"], [
    ["Một cây – thư viện", "19,6%", "9,9%", "0,842", "75%"],
    ["Một cây – tự cài đặt", "18,6%", "9,5%", "0,851", "75%"],
    ["**Rừng 300 cây – thư viện**", "**15,5%**", "**9,0%**", "**0,891**", "**80%**"],
    ["Rừng 100 cây – tự cài đặt", "15,7%", "9,0%", "0,887", "80%"],
], [5.4, 2.4, 2.4, 2.4, 2.8], caption="Một cây so với cả rừng trên tập test", bold_rows=(2,))

h2("4.6. Mô hình so sánh: Ridge Regression")
para("Ridge [1] dự đoán log giá bằng **tổ hợp tuyến tính** của các đặc trưng và tìm trọng số bằng cách tối thiểu hóa sai số bình phương cộng một **khoản phạt L2**:")
equation(r"J(w) = \frac{1}{2n}\sum_{i=1}^{n}(y_i - \hat y_i)^2 + \frac{\alpha}{2}\sum_{j=1}^{p} w_j^2, \qquad \hat y = w_0 + \sum_{j=1}^{p} w_j x_j")
para("Khoản phạt giữ các hệ số ổn định khi có nhiều cột one-hot tương quan (hãng Toyota và dòng Toyota Vios luôn đi cùng nhau). Dùng `sklearn.linear_model.Ridge`, chọn α trên lưới 25 giá trị từ 10⁻³ đến 10³ bằng **GridSearchCV + 5-fold CV** trên train: **α = 0,56**, CV RMSE (log) = 0,288.")
para("**Vai trò của Ridge trong đề tài:** là **mô hình tuyến tính tốt nhất có thể** trên cùng dữ liệu (đã có `age²`, `log(km)` để biểu diễn khấu hao cong, đã được tinh chỉnh như RF). Phần chênh lệch giữa Ridge và Random Forest chính là đóng góp của **quan hệ phi tuyến và tương tác** giữa các đặc trưng. Vì Ridge cùng tập train/test, cùng thước đo và cũng được tinh chỉnh bằng CV, so sánh là **công bằng**.")

h2("4.7. Kỹ thuật xử lý mất cân bằng")
para("Ba kỹ thuật được thử trên chính mô hình cuối (Random Forest thư viện, tham số tốt nhất), so với huấn luyện bình thường:")
table(["Kỹ thuật", "Cách làm (chỉ trong phần train của mỗi fold)"], [
    ["Không cân bằng", "Huấn luyện bình thường"],
    ["**Trọng số mẫu**", "Mẫu ở phân khúc hiếm có trọng số lớn hơn: 1/√(tần suất phân khúc), chuẩn hóa về trung bình 1"],
    ["**Random Oversampling**", "Nhân bản ngẫu nhiên mẫu của phân khúc hiếm đến tối thiểu 25% quy mô phân khúc lớn nhất"],
    ["**Random Undersampling**", "Bỏ bớt mẫu ở phân khúc đông, mỗi phân khúc giữ tối đa 8 lần số mẫu của phân khúc hiếm nhất"],
], [4.6, 11.4], caption="Các kỹ thuật xử lý mất cân bằng được thử nghiệm", align=["l", "l"])
para("**Cách đánh giá:** 5-fold CV phân tầng theo phân khúc giá **chỉ trên tập train**. Việc cân bằng chỉ áp dụng cho phần train của từng fold; phần validation **giữ nguyên phân phối thật**. Cân bằng trước khi chia sẽ làm bản sao của cùng một xe nằm ở cả hai phía và cho điểm cao giả tạo.")
para("**Vì sao dùng CV chứ không dùng OOB:** khi oversampling, bản sao của một xe có thể nằm trong mẫu bootstrap của một cây trong khi bản gốc là mẫu \"ngoài túi\" của chính cây đó. OOB khi đó đánh giá trên xe mà mô hình đã thấy và cho điểm đẹp giả tạo. CV tách validation **trước** khi nhân bản nên không gặp vấn đề này. Thí nghiệm cũng so sánh học trên giá gốc với học trên log giá. Tập test không được dùng ở bước này.")

# ================================================================== CHƯƠNG 5
chapter(5, "KẾT QUẢ VÀ ĐÁNH GIÁ")
h2("5.1. Kết quả tổng thể trên tập test")
para("Đây là lần **duy nhất** tập test được dùng để đánh giá. Mọi mô hình học trên cùng tập train (11.149 tin) và được đánh giá trên cùng 2.788 tin test.")
table(["Mô hình", "MAPE", "MedAPE", "MAE (tr)", "RMSE (tr)", "R² (log)", "±10%", "±20%", "Huấn luyện"], [
    ["Baseline (trung vị)", "83,5%", "40,0%", "337", "724", "0,00", "14%", "26%", "–"],
    ["Ridge – thư viện (so sánh)", "18,5%", "10,7%", "120", "411", "0,873", "48%", "75%", "0,04 s"],
    ["**Random Forest – thư viện**", "**15,5%**", "**9,0%**", "**102**", "**372**", "**0,891**", "54%", "**80%**", "9 s"],
    ["**Random Forest – tự cài đặt**", "15,7%", "9,0%", "106", "388", "0,887", "55%", "80%", "143 s"],
], [4.6, 1.5, 1.6, 1.4, 1.5, 1.5, 1.2, 1.2, 1.5], caption="Kết quả trên tập test", bold_rows=(2,), font=10.5)
figure("m06_so_sanh_tong.png", 16, "So sánh MAPE và tỉ lệ xe được định giá sát giữa các mô hình")
para("**Cách đọc bảng:**")
bullet("**Random Forest vượt xa mốc so sánh:** baseline sai trung bình 83,5% và R² gần 0; Random Forest giảm còn **15,5%**, R² (log) 0,891. Theo thang tham khảo ở Bảng 4.2, đây là mức **tốt** (10-20%).")
bullet("**MedAPE 9,0% thấp hơn hẳn MAPE 15,5%:** xe \"điển hình\" chỉ bị đoán lệch khoảng 9%; một nhóm nhỏ xe bị sai rất nặng (chủ yếu ở hai đầu thang giá, mục 5.2 và 5.4) kéo trung bình lên.")
bullet("**RMSE (372 triệu) gấp khoảng 3,6 lần MAE (102 triệu):** RMSE bình phương sai số nên phóng đại lỗi lớn; điều này cho thấy có một số xe bị đoán sai hàng tỷ đồng, chủ yếu là xe sang.")
bullet("**Lệch ≤ 20% đạt 80%:** 4 trên 5 xe được định giá sát trong khoảng ±20%; hơn một nửa sát trong ±10%. Ngay cả baseline đoán mù cũng trúng 26% vì nhiều xe có giá quanh mức trung vị, nên chỉ số này phải đọc cạnh baseline.")
bullet("**Bản tự cài đặt** gần như trùng bản thư viện về độ chính xác; khác biệt lớn nhất là thời gian huấn luyện (143 giây so với 9 giây).")

h2("5.2. Sai số theo phân khúc giá")
table(["Phân khúc", "% dữ liệu", "Baseline", "Ridge", "RF (thư viện)", "RF (tự cài đặt)"], [
    ["< 200 triệu", "14%", "349,8%", "43,0%", "34,7%", "35,6%"],
    ["200-500 triệu", "42%", "38,5%", "13,9%", "11,7%", "11,6%"],
    ["500-1.000 triệu", "33%", "30,7%", "12,6%", "10,9%", "10,8%"],
    ["1-2 tỷ", "8%", "64,1%", "16,7%", "15,7%", "15,8%"],
    ["> 2 tỷ", "3%", "85,4%", "34,3%", "27,0%", "30,4%"],
], [3.2, 2.0, 2.2, 2.2, 3.2, 3.2], caption="MAPE theo phân khúc giá trên tập test")
figure("m07_mape_phan_khuc.png", 15, "MAPE theo phân khúc giá")
bullet("**Phân khúc giữa (200 triệu - 1 tỷ, khoảng 75% thị trường):** sai số khoảng 11%, đủ tốt để tham khảo giá.")
bullet("**Xe dưới 200 triệu (35%):** giá thấp nên chênh vài chục triệu (tình trạng thực tế, thương lượng) đã là phần trăm lớn; xe đời cũ rất khác nhau về tình trạng mà dữ liệu không có; và còn sót tin ghi giá trả trước (mục 5.4).")
bullet("**Xe trên 2 tỷ (27%):** chỉ 3% dữ liệu và thiếu thông tin quyết định giá xe sang (phiên bản, trang bị).")
bullet("Random Forest tốt hơn Ridge ở **mọi phân khúc**. Baseline có sai số thấp nhất ở 500-1.000 triệu chỉ vì giá trung vị nằm sẵn trong khoảng này, cho thấy vì sao hai đầu thang giá luôn khó hơn.")

h2("5.3. Dự đoán so với giá thật và phần dư")
para("Trên Hình 5.3, mỗi điểm là một xe trong tập test; trục dùng thang log nên khoảng cách tới đường chéo chính là sai số phần trăm. Phần dư của cả hai mô hình đối xứng quanh 0 (không lệch hệ thống), Random Forest có phân bố hẹp hơn. Biểu đồ hộp cho thấy sai số tăng ở hai đầu thang giá.")
figure("m08_du_doan_vs_thuc_te.png", 16, "Giá dự đoán so với giá thật trên tập test (thang log)")
figure("m09_phan_du.png", 16, "Phân bố phần dư và sai số tuyệt đối theo phân khúc")

h2("5.4. Những xe bị đoán sai nhiều nhất")
para("Để hiểu sai số lớn đến từ đâu, xem cụ thể 10 xe Random Forest (thư viện) đoán lệch nhiều nhất trên tập test:")
table(["Hãng", "Dòng xe", "Năm", "Tình trạng", "Số km", "Giá thật", "Dự đoán", "Lệch"], [
    ["Hyundai", "Starex", "2023", "Đã dùng", "20.000", "35", "655", "+1.771%"],
    ["Isuzu", "Hi lander", "2023", "Đã dùng", "100.000", "60", "505", "+742%"],
    ["Toyota", "4 Runner", "2025", "Đã dùng", "9.000", "115", "819", "+612%"],
    ["Daewoo", "Matiz", "2016", "Đã dùng", "300.000", "23", "136", "+493%"],
    ["Hyundai", "Accent", "2026", "Mới", "0", "90", "464", "+416%"],
    ["Mitsubishi", "Xforce", "2025", "Đã dùng", "18.000", "170", "612", "+260%"],
    ["MINI", "Cooper", "2025", "Đã dùng", "91.000", "199", "612", "+208%"],
    ["Hyundai", "Tucson", "2021", "Đã dùng", "44.000", "190", "567", "+198%"],
    ["Ford", "Ranger", "2024", "Đã dùng", "36.750", "225", "616", "+174%"],
    ["Toyota", "Wigo", "2023", "Đã dùng", "12.000", "148", "379", "+157%"],
], [2.2, 2.2, 1.3, 2.0, 1.9, 2.0, 2.0, 2.4], caption="Mười xe bị đoán sai nhiều nhất trên tập test (giá: triệu đồng)", font=10.5)
para("**Nhận xét:** sai số lớn nhất phần nhiều đến từ **nhãn sai còn sót**, không phải từ mô hình. Phần lớn là xe **đời 2021-2026** nhưng rao chỉ 35-225 triệu (Hyundai Starex 2023 giá 35 triệu, Hyundai Accent 2026 **xe mới** giá 90 triệu): gần như chắc chắn là **giá trả trước / đặt cọc**. Luật lọc ở mục 3.5 chỉ áp dụng cho nhóm có ít nhất 3 xe cùng hãng-dòng-năm, nên các dòng xe ít tin lọt qua. Với những xe này, giá mô hình dự đoán (400-650 triệu) hợp lý hơn giá rao. Isuzu Hi lander đời \"2023\" nhiều khả năng là nhập sai năm sản xuất (dòng xe này đã ngừng bán từ lâu). Hệ quả: MAPE, nhất là phân khúc dưới 200 triệu, bị **đội lên** bởi các tin này.")

h2("5.5. Kiểm định ý nghĩa thống kê bằng bootstrap")
para("**Khoảng tin cậy là gì?** MAPE 15,5% được đo trên 2.788 xe cụ thể. Nếu tình cờ lấy 2.788 xe khác, con số có thể là 15,1% hay 16,0%. Khoảng tin cậy 95% cho biết con số này có thể dao động trong khoảng nào nếu đánh giá lại trên một tập dữ liệu khác.")
para("**Cách ước lượng bằng bootstrap [7]:** từ 2.788 xe test, bốc ngẫu nhiên 2.788 xe **có hoàn lại** để được một \"tập test giả\", tính lại MAPE; lặp lại 1.000 lần, được 1.000 giá trị MAPE; bỏ 2,5% thấp nhất và 2,5% cao nhất, khoảng còn lại là **khoảng tin cậy 95%**.")
para("**So sánh hai mô hình bằng bootstrap ghép cặp:** không nên chỉ nhìn hai khoảng tin cậy có chồng nhau không (cách này quá khắt khe). Cách đúng là trong mỗi lần bốc, tính **hiệu MAPE của hai mô hình trên cùng các xe đó**. Nếu khoảng tin cậy 95% của hiệu **không chứa 0**, chênh lệch có ý nghĩa thống kê; nếu **chứa 0**, chênh lệch có thể chỉ là ngẫu nhiên.")
table(["So sánh", "2,5%", "Trung vị", "97,5%", "Có ý nghĩa?"], [
    ["Ridge − Random Forest (thư viện)", "2,08", "2,97", "3,97", "**Có** (cả khoảng > 0)"],
    ["RF tự cài đặt − RF thư viện", "−0,04", "0,19", "0,46", "Không (khoảng chứa 0)"],
], [6.2, 1.8, 2.0, 1.8, 4.2], caption="Hiệu MAPE giữa các mô hình (điểm %), bootstrap ghép cặp 1.000 lần")
figure("m10_khoang_tin_cay.png", 13.5, "MAPE trên tập test với khoảng tin cậy 95% (bootstrap)")
para("**Kết luận:** Random Forest tốt hơn Ridge khoảng **3 điểm %** và chênh lệch **có ý nghĩa thống kê** (KTC 95%: 2,1-4,0). Rừng tự cài đặt và rừng thư viện **không khác nhau có ý nghĩa**.")

h2("5.6. Rừng tự cài đặt so với rừng thư viện trên từng xe")
para("Nếu hai bản cài đặt cùng một thuật toán, dự đoán của chúng phải gần như trùng nhau trên **từng xe**, không chỉ ở con số trung bình. Hình 5.6: các điểm bám sát đường chéo, độ lệch trung vị giữa hai bản là **1,7%**, và **83% số xe** lệch không quá 5%.")
figure("m11_tu_cai_dat_vs_thu_vien.png", 16, "Dự đoán của rừng tự cài đặt so với rừng thư viện trên từng xe trong tập test")
table(["Tiêu chí", "Thư viện", "Tự cài đặt"], [
    ["MAPE test", "15,5%", "15,7%"],
    ["OOB R² (log)", "0,896", "0,892"],
    ["Số cây / cách chọn tham số", "300 / CV", "100 / OOB"],
    ["Tìm ngưỡng chia", "chính xác, mọi giá trị", "histogram 32 bin"],
    ["Thứ hạng độ quan trọng đặc trưng", "tuổi, hộp số, hãng, dòng xe...", "giống hệt"],
    ["Thời gian huấn luyện", "9 giây (C/Cython, nhiều lõi)", "143 giây (Python, một lõi)"],
], [5.6, 5.2, 5.2], caption="Tóm tắt so sánh hai bản Random Forest", align=["l", "l", "l"], font=11)
para("Chênh lệch nhỏ giữa hai bản đến từ việc chia đặc trưng thành 32 bin, dùng 100 cây thay vì 300, và cấu hình tinh chỉnh bằng OOB hơi khác (`min_samples_leaf = 2`). Rõ nhất ở xe trên 2 tỷ (30,4% so với 27,0%), nơi ít mẫu nên bin thô và lá lớn hơn ảnh hưởng nhiều hơn.")

h2("5.7. Diễn giải mô hình")
h3("5.7.1. Độ quan trọng theo mức giảm sai số (impurity)")
para("Một đặc trưng phân loại bị one-hot thành nhiều cột, nên độ quan trọng của mọi cột thuộc cùng một đặc trưng gốc được cộng lại:")
figure("m12_do_quan_trong.png", 13, "Độ quan trọng đặc trưng (impurity): bản thư viện và bản tự cài đặt")
para("Tuổi xe đứng đầu (31%), sau đó là hộp số (18%), hãng (16%), dòng xe (13%). Hai bản cài đặt cho **cùng thứ hạng** và gần như cùng tỉ lệ, thêm một bằng chứng rừng tự viết hoạt động đúng. Tuy nhiên cách đo này có hai điểm yếu đã biết [10]: tính trên train (kể cả những phép chia chỉ học thuộc nhiễu), và **bị lệch khi các đặc trưng tương quan**. Ví dụ xe số sàn thường là xe đời cũ, nên hộp số và tuổi xe mang trùng thông tin; cây chọn đặc trưng nào để chia phần nhiều là ngẫu nhiên.")
h3("5.7.2. Kiểm tra lại bằng permutation importance")
para("**Permutation importance** [3] hỏi trực tiếp: nếu **xáo trộn ngẫu nhiên** một đặc trưng trên tập test (phá vỡ quan hệ của nó với giá, giữ nguyên phân phối), MAPE tăng bao nhiêu? Đo trên **đặc trưng gốc** (xáo trộn cả cột dòng xe rồi mới mã hóa), lặp 5 lần mỗi đặc trưng. Tập test chỉ dùng để diễn giải mô hình đã chọn, không để chọn gì.")
figure("m15_permutation_importance.png", 16, "Hai cách đo độ quan trọng đặc trưng của Random Forest")
table(["Đặc trưng", "Impurity (%)", "Permutation: MAPE tăng (điểm %)", "Hạng (impurity → permutation)"], [
    ["Tuổi xe", "31,0", "43,2", "1 → 1"],
    ["Hãng", "15,8", "19,7", "3 → 2"],
    ["Kiểu dáng", "7,4", "11,3", "5 → 3"],
    ["Hộp số", "17,7", "11,0", "**2 → 4**"],
    ["Dòng xe", "12,8", "8,8", "4 → 5"],
    ["Nhiên liệu", "3,1", "5,5", "7 → 6"],
    ["Số chỗ", "2,7", "4,6", "8 → 7"],
    ["Số km", "5,4", "3,7", "6 → 8"],
    ["Số km / năm", "2,5", "0,6", "9 → 9"],
    ["Xuất xứ, tình trạng, km là giá trị điền", "≤ 1,3", "≤ 0,2", "cuối"],
], [5.6, 2.6, 4.2, 3.6], caption="Độ quan trọng đặc trưng của Random Forest (thư viện) theo hai cách đo", font=11)
bullet("**Tuổi xe** quan trọng nhất theo cả hai cách: khấu hao là yếu tố lớn nhất quyết định giá xe cũ.")
bullet("**Hộp số tụt từ hạng 2 xuống hạng 4:** một phần độ quan trọng impurity của nó là thông tin **trùng với tuổi xe**. Khi xáo trộn hộp số, tuổi xe vẫn giữ phần thông tin chung đó.")
bullet("**Hãng lên hạng 2:** mức giá chung của thương hiệu là thông tin mà các đặc trưng khác không thay thế được. Dòng xe xếp sau hãng vì một phần thông tin đã có trong hãng, kiểu dáng, số chỗ.")
bullet("**Số km/năm, xuất xứ, tình trạng, cờ km là giá trị điền** gần như bằng 0: tình trạng mới/cũ đã nằm trong tuổi = 0 và km = 0.")
bullet("Xáo trộn từng đặc trưng độc lập có thể tạo tổ hợp không thực tế (xe 0 tuổi đi 200.000 km), nên con số dùng để **xếp hạng**, không phải tác động nhân quả.")
h3("5.7.3. Đường khấu hao")
para("Lấy ba mẫu xe phổ biến ở ba phân khúc, cố định mọi thuộc tính, chỉ thay đổi tuổi xe (giả định 15.000 km/năm; tuổi 0 là xe mới, 0 km), rồi xem hai mô hình dự đoán giá:")
figure("m13_duong_khau_hao.png", 16, "Giá dự đoán theo tuổi xe của ba mẫu xe ở ba phân khúc")
bullet("**Ridge** cho đường khấu hao trơn và **cùng một hình dạng** cho mọi mẫu xe (chỉ dịch lên xuống), vì mô hình tuyến tính trên log giá giả định mọi xe mất giá cùng một tỉ lệ % mỗi năm.")
bullet("**Random Forest** cho đường bậc thang với **hình dạng khác nhau giữa các dòng xe**: xe sang (Mercedes C Class) mất giá rất nhanh trong 3-4 năm đầu, khác hẳn xe phổ thông. Đây chính là **tương tác** giữa dòng xe và tuổi xe mà Ridge không biểu diễn được, và là lý do Random Forest chính xác hơn 3 điểm %.")

h2("5.8. Thí nghiệm log giá và xử lý mất cân bằng")
table(["Phương án", "Cỡ train/fold", "MAPE", "MedAPE", "< 200 tr", "200-500", "500-1.000", "1-2 tỷ", "> 2 tỷ"], [
    ["Giá gốc", "8.919", "19,6%", "9,7%", "51,3%", "15,8%", "11,5%", "15,6%", "22,5%"],
    ["**Log giá** (chọn)", "8.919", "16,2%", "9,1%", "35,7%", "12,8%", "10,8%", "16,5%", "27,4%"],
    ["Log + trọng số mẫu", "8.919", "16,3%", "8,8%", "34,9%", "13,6%", "11,2%", "15,3%", "24,1%"],
    ["Log + Oversampling", "9.830", "16,1%", "8,7%", "36,0%", "12,8%", "11,0%", "16,3%", "24,5%"],
    ["Log + Undersampling", "6.925", "16,8%", "9,5%", "35,5%", "14,3%", "11,0%", "16,0%", "26,7%"],
], [3.8, 1.8, 1.4, 1.5, 1.5, 1.5, 1.6, 1.4, 1.4], caption="Thí nghiệm trên Random Forest (5-fold CV trên tập train)", bold_rows=(1,), font=10.5)
figure("m14_log_va_mat_can_bang.png", 16, "MAPE tổng và MAPE theo phân khúc của các phương án (CV trên train)")
para("**Log giá tốt hơn rõ rệt.** MAPE tổng giảm từ 19,6% xuống 16,2%; phân khúc dưới 200 triệu cải thiện mạnh nhất (51% → 36%). Học trên giá gốc thì sai số bình phương bị xe đắt chi phối, nên mô hình bỏ qua sai số tương đối ở xe rẻ. Đổi lại, giá gốc tốt hơn ở xe trên 2 tỷ (22,5% so với 27,4%), cũng vì lý do đó.")
bullet("**MAPE tổng** của ba kỹ thuật cân bằng nằm trong khoảng 16,1-16,8%, so với 16,2% khi không cân bằng.")
bullet("**Trọng số mẫu** và **oversampling** giảm sai số ở xe trên 2 tỷ khoảng 3 điểm % (27,4% → 24,1% và 24,5%). Trọng số mẫu đổi lại bằng sai số cao hơn ở phân khúc đông nhất 200-500 triệu (12,8% → 13,6%); oversampling gần như không làm phân khúc nào kém đi.")
bullet("**Undersampling** bỏ khoảng 22% dữ liệu train, MAPE tổng tăng lên 16,8% và gần như không giúp xe sang: bỏ dữ liệu thật thì mất thông tin.")
bullet("Dù đã cân bằng, sai số ở xe trên 2 tỷ vẫn cao gấp đôi phân khúc giữa: nguyên nhân chủ yếu là **thiếu thông tin** (phiên bản, trang bị), không chỉ do ít mẫu. Nhân bản dữ liệu không tạo thêm thông tin mới.")
para("**Quyết định:** mô hình cuối **học trên log giá, không dùng kỹ thuật cân bằng**, vì thước đo chính không đổi và mô hình đơn giản hơn. Nếu ứng dụng ưu tiên định giá xe sang, nên dùng **Random Oversampling** (chỉ trên tập train).")

h2("5.9. Ảnh hưởng của tin đăng lại")
para("Bước tiền xử lý đã xóa các tin trùng hoàn toàn, nhưng cùng một chiếc xe có thể được **đăng lại với giá chênh nhẹ**. Khi chia ngẫu nhiên, một bản rơi vào train và bản kia vào test; Random Forest với lá nhỏ chỉ cần \"nhớ\" giá là đoán đúng. Hai cách kiểm tra:")
para("**(1) Trên tập test:** tin test nào có một tin train **giống hệt mọi đặc trưng, kể cả số km**?", indent=False)
table(["Nhóm tin test", "Số tin", "% test", "MAPE Ridge", "MAPE RF (TV)", "MAPE RF (TCĐ)"], [
    ["Có bản sao trong train", "391", "14,0%", "14,0%", "10,9%", "10,8%"],
    ["   trong đó xe đã dùng", "247", "8,9%", "14,0%", "10,5%", "10,7%"],
    ["Không có bản sao", "2.397", "86,0%", "19,2%", "16,2%", "16,5%"],
    ["Toàn bộ test", "2.788", "100%", "18,5%", "15,5%", "15,7%"],
], [4.8, 1.8, 1.8, 2.4, 2.6, 2.6], caption="Sai số trên nhóm tin có và không có bản sao trong tập train", font=11)
para("**(2) CV chia theo nhóm trên tập train:** chạy lại 5-fold CV cho Random Forest theo hai cách: chia ngẫu nhiên (`StratifiedKFold`) và chia theo nhóm (`StratifiedGroupKFold`), trong đó mọi tin giống hệt nhau về đặc trưng luôn nằm **cùng một phía**. Train có 10.175 nhóm trên 11.149 tin.", indent=False)
table(["Cách chia", "MAPE", "< 200 tr", "200-500", "500-1.000", "1-2 tỷ", "> 2 tỷ"], [
    ["Ngẫu nhiên (StratifiedKFold)", "16,1%", "35,7%", "12,8%", "10,8%", "16,5%", "27,6%"],
    ["Theo nhóm (StratifiedGroupKFold)", "16,7%", "38,0%", "13,0%", "10,9%", "16,9%", "28,7%"],
], [5.6, 1.6, 1.7, 1.7, 1.8, 1.6, 1.6], caption="MAPE (CV trên train) khi chia ngẫu nhiên và chia theo nhóm", font=11)
para("Chia theo nhóm làm MAPE tăng khoảng **0,5 điểm %**, rõ nhất ở xe dưới 200 triệu. Mức lạc quan do tin đăng lại là **có nhưng nhỏ** và không đổi kết luận nào. **Ước lượng thận trọng** cho một chiếc xe mô hình chưa từng thấy: MAPE khoảng **16-16,5%**.")

h2("5.10. Lưu và sử dụng mô hình")
table(["File", "Nội dung"], [
    ["`preprocessor_rf.joblib`, `preprocessor_ridge.joblib`", "Bộ mã hóa của từng mô hình (fit trên train)"],
    ["`random_forest_sklearn.joblib`, `ridge_sklearn.joblib`", "Mô hình thư viện"],
    ["`random_forest_scratch.npz`", "Các cây và ngưỡng bin của rừng tự cài đặt (mảng NumPy, không phụ thuộc lớp trong notebook)"],
    ["`metadata.json`", "Danh sách đặc trưng, bảng gộp nhãn, trung vị số km theo năm, tham số tốt nhất, metric"],
    ["`price_model.py`", "Hàm `load_model` / `predict_price`: nhận thông tin thô của xe, tự dựng đặc trưng như lúc huấn luyện"],
], [7.0, 9.0], caption="Các file mô hình đã lưu trong thư mục model/", align=["l", "l"], font=11)
para("Khi nạp lại bằng `price_model.py` và truyền thông tin thô của các xe test, dự đoán trùng với notebook (lệch lớn nhất khoảng 10⁻¹¹ triệu). Ví dụ:")
table(["Xe", "RF (thư viện)", "RF (tự cài đặt)", "Ridge"], [
    ["Toyota Vios 2018, 80.000 km", "303", "303", "281"],
    ["Ford Ranger 2021, không rõ số km", "515", "536", "596"],
], [6.4, 3.2, 3.2, 3.2], caption="Ví dụ dự đoán giá (triệu đồng)")

# ================================================================== CHƯƠNG 6
chapter(6, "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
h2("6.1. Kết quả đạt được")
bullet("Xây dựng hoàn chỉnh quy trình học máy từ đầu đến cuối: thu thập **14.835 tin** từ Chợ Tốt, tiền xử lý còn **13.937 tin** sạch, huấn luyện, đánh giá, diễn giải và đóng gói mô hình để sử dụng.")
bullet("**Random Forest** đạt MAPE **15,5%**, MedAPE 9,0%, R² (log) 0,891, **80% xe** được định giá lệch không quá 20%; tốt hơn Ridge (18,5%) khoảng 3 điểm %, **có ý nghĩa thống kê** (KTC 95%: 2,1-4,0). Phần chênh này là đóng góp của quan hệ phi tuyến và tương tác giữa các đặc trưng.")
bullet("**Random Forest tự cài đặt bằng NumPy** (CART tìm ngưỡng bằng histogram, bootstrap, chọn ngẫu nhiên đặc trưng, OOB, độ quan trọng đặc trưng) được kiểm chứng ở cả mức một cây và cả rừng: MAPE 15,7%, khác bản thư viện không có ý nghĩa thống kê, dự đoán trên từng xe lệch trung vị 1,7%.")
bullet("Mọi quyết định đều kiểm chứng bằng CV/OOB trên train: one-hot tốt hơn mã số thứ tự; học trên log giá giảm MAPE từ 19,6% xuống 16,2%; trọng số mẫu và oversampling giảm khoảng 3 điểm % sai số xe trên 2 tỷ; undersampling làm kết quả kém đi.")
bullet("Yếu tố quyết định giá (permutation importance): tuổi xe, hãng, kiểu dáng, hộp số, dòng xe.")

h2("6.2. Hạn chế")
bullet("**Giá rao bán, không phải giá giao dịch:** luôn có phần thương lượng mà mô hình không biết.")
bullet("**Thiếu đặc trưng quan trọng:** trang tin có các mục **số đời chủ, còn hạn đăng kiểm, phụ kiện đi kèm** và tiêu đề tin thường ghi **phiên bản xe**, nhưng code thu thập chưa lấy. Đây là nguyên nhân chính của sai số cao ở xe trên 2 tỷ (27%).")
bullet("**Nhãn sai còn sót:** luật lọc giá trả trước chỉ áp dụng cho nhóm ≥ 3 xe, nên vẫn còn tin ghi giá trả trước (mục 5.4), làm MAPE phân khúc dưới 200 triệu bị đội lên. Mặt khác, luật này dùng giá trên toàn bộ dữ liệu (là bước sửa nhãn), nên khi dùng thật mô hình vẫn sẽ gặp các tin như vậy.")
bullet("**Tin đăng lại:** 14% tin test có bản sao trong train; ước lượng thận trọng khoảng 16-16,5% thay vì 15,5%.")
bullet("**Mẫu dữ liệu:** chỉ lấy 750 trang tin mới nhất tại một thời điểm (10/2026), nghiêng về tin đăng gần đây và đại lý đăng nhiều; xe trên 10 tỷ đã bị loại nên mô hình không áp dụng cho siêu xe.")
bullet("**Rò rỉ nhẹ trong tiền xử lý:** điền thiếu và đếm nhãn hiếm dùng thống kê trên toàn bộ dữ liệu trước khi chia (chỉ dựa trên đặc trưng, không dùng giá). Rủi ro thấp nhưng cần nêu rõ.")

h2("6.3. Hướng phát triển")
bullet("Crawl thêm **tiêu đề, mô tả** và các mục số đời chủ, đăng kiểm: trích phiên bản xe, lọc giá trả trước theo từ khóa (\"trả trước\", \"góp\", \"vay\", \"cọc\").")
bullet("Chia train/test theo nhóm (`StratifiedGroupKFold`) làm cách đánh giá chính.")
bullet("Thử **Gradient Boosting** (học tuần tự các cây sửa lỗi của nhau), có thể dùng lại lớp cây đã tự cài đặt.")
bullet("Sau khi đánh giá, huấn luyện lại trên toàn bộ dữ liệu trước khi triển khai, và huấn luyện lại định kỳ khi thị trường thay đổi.")
bullet("Xây dựng giao diện web nhỏ cho phép nhập thông tin xe và nhận giá tham khảo, dựa trên `price_model.py`.")

# ================================================================== TÀI LIỆU THAM KHẢO
page_break()
CHAPTER[0] = 7
h1("TÀI LIỆU THAM KHẢO")
refs = [
    "Hoerl, A. E., & Kennard, R. W. (1970). Ridge regression: Biased estimation for nonorthogonal problems. *Technometrics*, 12(1), 55–67.",
    "Breiman, L., Friedman, J., Olshen, R., & Stone, C. (1984). *Classification and Regression Trees*. Wadsworth.",
    "Breiman, L. (2001). Random forests. *Machine Learning*, 45(1), 5–32.",
    "Hastie, T., Tibshirani, R., & Friedman, J. (2009). *The Elements of Statistical Learning* (2nd ed.). Springer.",
    "Ke, G., Meng, Q., Finley, T., et al. (2017). LightGBM: A highly efficient gradient boosting decision tree. *Advances in Neural Information Processing Systems 30*.",
    "Pedregosa, F., et al. (2011). Scikit-learn: Machine learning in Python. *Journal of Machine Learning Research*, 12, 2825–2830.",
    "Efron, B., & Tibshirani, R. J. (1993). *An Introduction to the Bootstrap*. Chapman & Hall.",
    "Chợ Tốt Xe. Chuyên mục mua bán ô tô. https://xe.chotot.com/mua-ban-oto (truy cập 25/10/2026).",
    "Lewis, C. D. (1982). *Industrial and Business Forecasting Methods*. Butterworths.",
    "Strobl, C., Boulesteix, A.-L., Zeileis, A., & Hothorn, T. (2007). Bias in random forest variable importance measures: Illustrations, sources and a solution. *BMC Bioinformatics*, 8, 25.",
]
for k, r in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.first_line_indent = Cm(-1)
    add_runs(p, f"[{k}]\t" + r)

# ================================================================== PHỤ LỤC
page_break()
CHAPTER[0] = "PL"
h1("PHỤ LỤC")
h2("A. Cấu trúc mã nguồn")
table(["Đường dẫn", "Nội dung"], [
    ["`src/crawl_data.ipynb`", "Thu thập dữ liệu từ Chợ Tốt"],
    ["`src/visualize_raw_data.ipynb`", "Biểu đồ khám phá dữ liệu thô"],
    ["`src/preprocess_data.ipynb`", "Tiền xử lý, tạo `clean_data_oto.csv`"],
    ["`src/model.ipynb`", "Random Forest (thư viện + tự cài đặt), Ridge để so sánh, đánh giá, thí nghiệm, xuất mô hình"],
    ["`src/figures/`", "Các biểu đồ dùng trong báo cáo"],
    ["`model/price_model.py`", "Nạp mô hình và dự đoán giá từ thông tin thô"],
    ["`slides/`", "Slide trình bày"],
], [6.0, 10.0], caption="Cấu trúc thư mục dự án", align=["l", "l"])
h2("B. Cách chạy")
para("Chạy các notebook với thư mục làm việc là `src` theo thứ tự: `crawl_data` (không bắt buộc, đã có dữ liệu), `preprocess_data` (khoảng 2 phút), `model` (khoảng 15 phút, phần lớn là rừng tự cài đặt). Mọi bước ngẫu nhiên cố định `random_state = 42` nên kết quả tái lập được. Dùng mô hình đã lưu:", indent=False)
code = ['from price_model import load_model, predict_price', 'import pandas as pd',
        'xe = pd.DataFrame([{"carbrand": "Toyota", "carmodel": "Vios", "carorigin": "Việt Nam",',
        '    "condition_ad": "Đã sử dụng", "gearbox": "Tự động", "fuel": "Xăng",',
        '    "cartype": "Sedan", "carseats": 5, "mfdate": 2018, "mileage_km": 80000}])',
        'print(predict_price(load_model("random_forest_sklearn"), xe))   # ~303 triệu']
for line in code:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(9.5)
doc.add_paragraph()
h2("C. Phân công công việc")
table(["Thành viên", "Công việc"], [
    ["<Họ tên SV 1>", "<Nội dung công việc>"],
    ["<Họ tên SV 2>", "<Nội dung công việc>"],
], [5.0, 11.0], caption="Phân công công việc trong nhóm", align=["l", "l"])

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print("Saved", OUT)
